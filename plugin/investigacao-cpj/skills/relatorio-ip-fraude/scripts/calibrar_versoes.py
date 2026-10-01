#!/usr/bin/env python3
"""K01 — compara localmente a minuta gerada com o FINAL do investigador e propõe lições genéricas.

Nada sai do computador. O comparativo usa só contagens (palavras, citações, termos de cautela, tópicos...), nunca
copia trecho, nome ou número do caso. As lições vêm de um catálogo fixo (CATALOGO): é impossível um dado de caso
chegar a calibracao\\ por aqui. Só entra no arquivo o que o investigador aprova, por id.

Uso (CLI): calibrar_versoes.py <minuta.md> <final.(md|docx|pdf)>   -> imprime a proposta em JSON (não grava)
"""
import datetime
import difflib
import io
import json
import os
import re
import sys

SECAO = "## 8. Lições aprovadas pelo comparativo minuta × FINAL"

CATALOGO = {
    "mais-conciso": "O investigador costuma enxugar o texto da minuta: redigir mais conciso, sem repetir fatos já ditos.",
    "mais-completo": "O investigador costuma acrescentar conteúdo à minuta: detalhar mais a dinâmica dos fatos e as circunstâncias.",
    "reforcar-cautela": "Reforçar a linguagem prudente (\"em tese\", \"investigado(a)\", \"há indícios\") ao tratar de autoria.",
    "menos-citacoes": "Reduzir a densidade de citações de página/folha: manter apenas nos pontos-chave e nos valores.",
    "mais-citacoes": "Ampliar as citações de página/folha: o investigador acrescenta a origem em mais afirmações.",
    "sem-topicos": "Preferir texto corrido; o investigador converte tópicos e listas em parágrafos.",
    "paragrafos-curtos": "Preferir parágrafos mais curtos e diretos.",
    "paragrafos-longos": "Agrupar ideias em parágrafos mais completos, em vez de frases soltas.",
    "menos-valores": "Citar menos valores de forma dispersa: concentrar os valores no resumo da dinâmica financeira.",
    "mais-valores": "Detalhar mais os valores e datas das transferências no texto.",
    "limpar-pendencias": "Eliminar marcadores de pendência e texto-guia do modelo antes de entregar a minuta.",
    "reescrita-profunda": "A minuta foi muito reescrita: conferir aderência ao estilo CPJ (texto corrido, foco na dinâmica e nos valores).",
}

RE_CAUTELA = re.compile(r"\bem tese\b|\binvestigad[oa]s?\b|\bind[íi]cios?\b|\bsupost[oa]s?\b|\bsupostamente\b", re.I)
RE_CITA = re.compile(r"\b(?:p[áa]gs?\.?|p[áa]ginas?|fls?\.?)\s*\d", re.I)
RE_VALOR = re.compile(r"R\$\s*[\d.]+,\d{2}")
RE_TOPICO = re.compile(r"^\s*(?:[-*•–]|\d{1,2}[.)])\s+\S", re.M)
RE_PENDENTE = re.compile(r"\{[^{}\n]{1,200}\}|\[\s*PENDENTE[^\]\n]*\]|\bPREENCHER\b|\(A\)|A\(o\)", re.I)


def ler_texto(nome, dados):
    """Texto do arquivo enviado (md/txt/docx/pdf), lido em memória."""
    ext = os.path.splitext(nome or "")[1].lower()
    if ext in (".md", ".txt", ""):
        return dados.decode("utf-8", errors="replace")
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(dados))
        partes = [p.text for p in d.paragraphs]
        for t in d.tables:
            partes += [c.text for r in t.rows for c in r.cells]
        return "\n".join(partes)
    if ext == ".pdf":
        import pypdf
        r = pypdf.PdfReader(io.BytesIO(dados))
        return "\n".join((p.extract_text() or "") for p in r.pages)
    raise ValueError("Formato não suportado: envie PDF, DOCX ou MD.")


def _corpo(txt):
    """Remove o cabeçalho YAML da minuta, se houver."""
    if txt.startswith("---"):
        fim = txt.find("\n---", 3)
        if fim > 0:
            return txt[fim + 4:]
    return txt


def metricas(txt):
    txt = _corpo(txt)
    palavras = re.findall(r"\w+", txt)
    paragrafos = [p for p in re.split(r"\n\s*\n", txt) if len(p.split()) >= 5 and not RE_TOPICO.match(p)]
    n = max(len(palavras), 1)
    return {
        "palavras": len(palavras),
        "cautela_por_mil": round(1000 * len(RE_CAUTELA.findall(txt)) / n, 1),
        "citacoes_por_mil": round(1000 * len(RE_CITA.findall(txt)) / n, 1),
        "valores": len(RE_VALOR.findall(txt)),
        "topicos": len(RE_TOPICO.findall(txt)),
        "pendencias": len(RE_PENDENTE.findall(txt)),
        "palavras_por_paragrafo": round(sum(len(p.split()) for p in paragrafos) / max(len(paragrafos), 1), 1),
    }


def _sentencas(txt):
    return [re.sub(r"\s+", " ", s).strip().lower() for s in re.split(r"(?<=[.!?])\s+", _corpo(txt)) if len(s.split()) >= 4]


def similaridade(minuta, final):
    a, b = _sentencas(minuta), _sentencas(final)
    if not a or not b:
        return 0.0
    conj = set(b)
    return round(sum(1 for s in a if s in conj or difflib.get_close_matches(s, b, 1, 0.85)) / len(a), 2)


def comparar(minuta, final):
    """{'metricas': {...}, 'licoes': [{'id', 'texto'}]} — só números e frases do catálogo."""
    m, f = metricas(minuta), metricas(final)
    ids = []
    if m["palavras"] and f["palavras"] < 0.8 * m["palavras"]:
        ids.append("mais-conciso")
    elif m["palavras"] and f["palavras"] > 1.2 * m["palavras"]:
        ids.append("mais-completo")
    if f["cautela_por_mil"] >= 1.5 * max(m["cautela_por_mil"], 0.5) and f["cautela_por_mil"] - m["cautela_por_mil"] >= 1:
        ids.append("reforcar-cautela")
    if m["citacoes_por_mil"] and f["citacoes_por_mil"] < 0.7 * m["citacoes_por_mil"]:
        ids.append("menos-citacoes")
    elif f["citacoes_por_mil"] > 1.3 * max(m["citacoes_por_mil"], 1):
        ids.append("mais-citacoes")
    if m["topicos"] >= 3 and f["topicos"] <= m["topicos"] // 2:
        ids.append("sem-topicos")
    if m["palavras_por_paragrafo"] and f["palavras_por_paragrafo"] < 0.75 * m["palavras_por_paragrafo"]:
        ids.append("paragrafos-curtos")
    elif m["palavras_por_paragrafo"] and f["palavras_por_paragrafo"] > 1.3 * m["palavras_por_paragrafo"]:
        ids.append("paragrafos-longos")
    if m["valores"] >= 6 and f["valores"] < 0.6 * m["valores"]:
        ids.append("menos-valores")
    elif m["valores"] and f["valores"] > 1.5 * m["valores"]:
        ids.append("mais-valores")
    if m["pendencias"] and not f["pendencias"]:
        ids.append("limpar-pendencias")
    sim = similaridade(minuta, final)
    if sim < 0.4:
        ids.append("reescrita-profunda")
    return {"metricas": {"minuta": m, "final": f, "similaridade": sim},
            "licoes": [{"id": i, "texto": CATALOGO[i]} for i in ids]}


def registrar(ws, ids, data=None):
    """Grava em calibracao\\ somente lições do catálogo, escolhidas por id. Devolve (gravadas, ja_existentes)."""
    validos = [i for i in dict.fromkeys(ids) if i in CATALOGO]
    if not validos:
        raise ValueError("Nenhuma lição válida para gravar.")
    data = data or datetime.date.today().isoformat()
    arq = os.path.join(ws, "calibracao", "licoes-aprendidas.md")
    with open(arq, encoding="utf-8") as f:
        txt = f.read()
    novas = [i for i in validos if CATALOGO[i] not in txt]
    if novas:
        linhas = "".join(f"- **{data} · calibração aprovada pelo investigador:** {CATALOGO[i]}\n" for i in novas)
        if SECAO in txt:
            txt = txt.rstrip("\n") + "\n" + linhas
        else:
            txt = txt.rstrip("\n") + f"\n\n---\n\n{SECAO}\n\n" + linhas
        _gravar(arq, txt)
        hist = os.path.join(ws, "calibracao", "historico-calibracao.md")
        with open(hist, encoding="utf-8") as f:
            h = f.read().rstrip("\n")
        _gravar(hist, h + f"\n| {data} | — | comparativo minuta × FINAL (contagens, sem dados de caso) | {len(novas)} | "
                          f"{len(validos) - len(novas)} já existente(s) | — |\n")
    return novas, [i for i in validos if i not in novas]


def _gravar(caminho, texto):
    tmp = caminho + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="\n") as f:
        f.write(texto)
    os.replace(tmp, caminho)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as f:
        mi = f.read()
    with open(sys.argv[2], "rb") as f:
        fi = ler_texto(sys.argv[2], f.read())
    print(json.dumps(comparar(mi, fi), ensure_ascii=False, indent=2))
