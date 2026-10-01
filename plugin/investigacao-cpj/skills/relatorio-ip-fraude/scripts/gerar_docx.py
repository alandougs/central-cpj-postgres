#!/usr/bin/env python3
"""Preenche o MODELO RELATÓRIO DE INVESTIGAÇÃO - CPJ (DOCX com timbre, rodapé e assinatura)
a partir de uma minuta Markdown.

Uso:
  gerar_docx.py minuta.md --saida RELATORIO.docx [--modelo modelo.docx] [--sem-assinatura]

Minuta (UTF-8):
  ---
  ordem_servico: 123/2026
  referencia: IPe nº 1234567-89.2026 / Processo nº 1500123-45.2026.8.26.0482 # DETERMINAÇÃO: SOMENTE IPe e Processo; sem BO nem IP local
  natureza: Estelionato (art. 171, CP)
  investigados: FULANO DE TAL
  vitimas: BELTRANA DE TAL
  local: Rua X, 100, Presidente Prudente/SP
  data_fatos: 10/03/2026
  local_data: Presidente Prudente, SP, 27 de setembro de 2026
  data_rodape: 27/09/2026
  delegado: NOME DO DELEGADO (sem "Dr.": o tratamento sai de delegado_genero)
  delegado_genero: M          # M ou F — concordância de gênero na saudação e no endereçamento final
  ---
  ## RESUMO DOS FATOS
  ...
  ## DILIGÊNCIAS REALIZADAS
  ...
  ## CONCLUSÃO
  ...

Markdown aceito no corpo: parágrafos, **negrito**, `### subtítulo` (negrito), listas `- ` / `1. `,
tabelas `| a | b |` (viram tabela Word), `<!-- notas internas -->` (removidas).
Campo ausente/vazio na minuta: o texto-guia do modelo é mantido para preenchimento manual.
Timbre (cabeçalho), rodapé com paginação automática e imagem de assinatura são preservados.
"""
import argparse, copy, os, re, sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

def ws_padrao():
    """Workspace CPJ: CPJ_WORKSPACE > 1ª pasta acima deste script com casos/ e modelos/ > cwd com casos/ > erro."""
    if os.environ.get("CPJ_WORKSPACE"):
        return os.environ["CPJ_WORKSPACE"]
    d = os.path.dirname(os.path.abspath(__file__))
    while True:
        if os.path.isdir(os.path.join(d, "casos")) and os.path.isdir(os.path.join(d, "modelos")):
            return d
        if os.path.dirname(d) == d:
            break
        d = os.path.dirname(d)
    if os.path.isdir(os.path.join(os.getcwd(), "casos")):
        return os.getcwd()
    raise SystemExit("Workspace CPJ não encontrado: defina CPJ_WORKSPACE ou execute a partir da pasta do workspace "
                     "(a que contém as pastas casos e modelos).")


def _resolver_modelo_padrao():
    try:
        ws = ws_padrao()
        return os.path.join(ws, "modelos", "MODELO RELATORIO DE INVESTIGACAO - CPJ 2026.docx")
    except SystemExit:
        return None


MODELO_PADRAO = _resolver_modelo_padrao()
FONTE, TAM = "Arial", 12

CAMPOS_CABECALHO = [  # (chave da minuta, rótulo no modelo)
    ("ordem_servico", "Ordem de Serviço:"),
    ("referencia", "Referência:"),
    ("natureza", "Natureza:"),
    ("investigados", "Investigado (s):"),
    ("vitimas", "Vítima(s):"),
    ("local", "Local:"),
    ("data_fatos", "Data dos Fatos:"),
    ("escrivao", "Escrivão do feito:"),
]
SECOES = {  # seção canônica da minuta -> prefixos dos parágrafos-guia do modelo que ela substitui
    "RESUMO DOS FATOS": ["{Resumir"],
    "DILIGÊNCIAS REALIZADAS": ["{Elencar"],
    "CONCLUSÃO": ["{breve descrição", "{Sugestão de providências"],
}
SECOES_OBRIGATORIAS = ["RESUMO DOS FATOS", "DILIGÊNCIAS REALIZADAS", "CONCLUSÃO"]


def remover_acentos(texto: str) -> str:
    import unicodedata
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def normalizar_titulo(texto: str) -> str:
    sem_acento = remover_acentos(texto).strip()
    limpo = re.sub(r"^\d+[\.\)\-\s]+", "", sem_acento)
    return limpo.strip().upper()


SECOES_CANONICAS = {normalizar_titulo(k): k for k in SECOES}

ap = argparse.ArgumentParser()
ap.add_argument("minuta"); ap.add_argument("--saida", required=True)
ap.add_argument("--modelo", default=None, help="caminho do arquivo modelo.docx (padrão: <workspace>/modelos/MODELO RELATORIO DE INVESTIGACAO - CPJ 2026.docx)")
ap.add_argument("--permitir-secoes-faltantes", action="store_true", help="não aborta se faltar seção obrigatória")
ap.add_argument("--sem-assinatura", action="store_true", help="remove a imagem da assinatura (rascunho)")
ap.add_argument("--fluxo-csv", default=None, help="caminho do arquivo fluxo-financeiro.csv para gerar fluxograma")
ap.add_argument("--fluxograma", default=None, help="caminho de imagem PNG existente do fluxograma")
ap.add_argument("--sem-fluxograma", action="store_true", help="não injeta fluxograma automaticamente")
a = ap.parse_args()


# ---------- leitura da minuta ----------
bruto = open(a.minuta, encoding="utf-8-sig").read()
bruto = re.sub(r"<!--.*?-->", "", bruto, flags=re.S)
meta, corpo = {}, bruto
m = re.match(r"\s*---\s*\n(.*?)\n---\s*\n(.*)", bruto, flags=re.S)
if m:
    for ln in m.group(1).splitlines():
        if ":" in ln:
            k, v = ln.split(":", 1)
            meta[k.strip().lower()] = v.strip()
    corpo = m.group(2)

secoes = {s: [] for s in SECOES}
atual = None
for ln in corpo.splitlines():
    m_head = re.match(r"^(#{2,3})\s+(.+?)\s*$", ln)
    if m_head:
        nivel = len(m_head.group(1))
        tit_raw = m_head.group(2).strip()
        tit_norm = normalizar_titulo(tit_raw)

        canon_match = None
        for k_norm, k_canon in SECOES_CANONICAS.items():
            if tit_norm == k_norm or k_norm in tit_norm or (len(tit_norm) >= 6 and tit_norm in k_norm):
                canon_match = k_canon
                break

        if canon_match:
            atual = canon_match
            continue
        elif nivel == 2:
            atual = tit_raw.upper()
            if atual not in secoes:
                secoes[atual] = []
            continue
        else:
            # Subtítulo (###) dentro da seção atual
            if atual:
                secoes[atual].append(ln)
            continue

    if atual:
        secoes[atual].append(ln)

faltantes = [s for s in SECOES_OBRIGATORIAS if not any(x.strip() for x in secoes.get(s, []))]
if faltantes and not a.permitir_secoes_faltantes:
    raise SystemExit(f"Erro: seção(ões) obrigatória(s) ausente(s) ou vazia(s) na minuta: {', '.join(faltantes)}")


def blocos(linhas):
    """Agrupa linhas em blocos: ('p', texto) | ('h', texto) | ('li', texto) | ('tab', [[...]])."""
    out, par, tab = [], [], []

    def fecha_par():
        if par: out.append(("p", " ".join(x.strip() for x in par))); par.clear()

    def fecha_tab():
        if tab:
            linhas_t = []
            alinhamentos = []
            for r in tab:
                cells = [c.strip() for c in r.strip().strip("|").split("|")]
                if re.fullmatch(r"\|?[\s:\-|]+\|?", r.strip()):
                    alinhamentos = []
                    for c in cells:
                        cs = c.strip()
                        if cs.startswith(":") and cs.endswith(":"):
                            alinhamentos.append(WD_ALIGN_PARAGRAPH.CENTER)
                        elif cs.endswith(":"):
                            alinhamentos.append(WD_ALIGN_PARAGRAPH.RIGHT)
                        elif cs.startswith(":"):
                            alinhamentos.append(WD_ALIGN_PARAGRAPH.LEFT)
                        else:
                            alinhamentos.append(None)
                else:
                    linhas_t.append(cells)
            if linhas_t:
                out.append(("tab", {"linhas": linhas_t, "alinhamentos": alinhamentos}))
            tab.clear()

    for ln in linhas:
        s = ln.strip()
        if s.startswith("|"):
            fecha_par(); tab.append(s); continue
        fecha_tab()
        if not s:
            fecha_par(); continue
        if s.startswith("###"):
            fecha_par(); out.append(("h", s.lstrip("#").strip())); continue
        if re.match(r"^([-*•]|\d+[.)])\s+", s):
            fecha_par(); out.append(("li", s)); continue
        if s.startswith("![") and "](" in s and s.endswith(")"):
            fecha_par()
            m_img = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
            if m_img:
                out.append(("img", (m_img.group(1).strip(), m_img.group(2).strip())))
            continue
        par.append(s)
    fecha_par(); fecha_tab()
    return out


# ---------- utilitários DOCX ----------
def parse_inline_tokens(texto, negrito_base=False):
    """Interpreta formatações ricas do Markdown:
    - Nomes/termos em negrito: **texto**
    - Informações super relevantes: ==grifo amarelo== ou <mark>grifo amarelo</mark>
    - Informações a preencher/obter manualmente: [PESQUISAR...], [OBTER...], [PREENCHER...], {PREENCHER...} -> CAIXA ALTA E VERMELHO
    """
    padrao = re.compile(
        r'('
        r'\[(?:PESQUISAR|OBTER|PREENCHER|LOCALIZAR|DILIG[EÊ]NCIA|HUMANO|ALERTA|VERIFICAR)[^\]]*\]|'
        r'\{(?:PESQUISAR|OBTER|PREENCHER|LOCALIZAR|DILIG[EÊ]NCIA|HUMANO|ALERTA|VERIFICAR)[^\}]*\}|'
        r'\[VERMELHO:[^\]]+\]|'
        r'==.+?==|'
        r'<mark>.+?</mark>|'
        r'\*\*.+?\*\*'
        r')',
        flags=re.IGNORECASE
    )
    pos = 0
    tokens = []
    for m in padrao.finditer(texto):
        start, end = m.span()
        if start > pos:
            tokens.append({'texto': texto[pos:start], 'bold': negrito_base, 'yellow': False, 'red': False})
        
        trecho = m.group(0)
        if trecho.startswith('==') and trecho.endswith('=='):
            conteudo = trecho[2:-2]
            is_bold = negrito_base or ('**' in conteudo)
            tokens.append({'texto': conteudo.replace('**', ''), 'bold': is_bold, 'yellow': True, 'red': False})
        elif trecho.lower().startswith('<mark>') and trecho.lower().endswith('</mark>'):
            conteudo = trecho[6:-7]
            is_bold = negrito_base or ('**' in conteudo)
            tokens.append({'texto': conteudo.replace('**', ''), 'bold': is_bold, 'yellow': True, 'red': False})
        elif trecho.startswith('**') and trecho.endswith('**'):
            conteudo = trecho[2:-2]
            if conteudo.startswith('==') and conteudo.endswith('=='):
                tokens.append({'texto': conteudo[2:-2], 'bold': True, 'yellow': True, 'red': False})
            elif re.search(r'^(?:\[|\{)(?:PESQUISAR|OBTER|PREENCHER|LOCALIZAR|DILIG[EÊ]NCIA|HUMANO|ALERTA|VERIFICAR)', conteudo, re.IGNORECASE):
                tokens.append({'texto': conteudo.upper(), 'bold': True, 'yellow': False, 'red': True})
            else:
                tokens.append({'texto': conteudo, 'bold': True, 'yellow': False, 'red': False})
        elif trecho.startswith('[VERMELHO:') and trecho.endswith(']'):
            conteudo = trecho[10:-1].strip()
            tokens.append({'texto': conteudo.upper(), 'bold': True, 'yellow': False, 'red': True})
        else:
            tokens.append({'texto': trecho.upper(), 'bold': True, 'yellow': False, 'red': True})
        pos = end
        
    if pos < len(texto):
        tokens.append({'texto': texto[pos:], 'bold': negrito_base, 'yellow': False, 'red': False})
    return tokens


def add_runs(p, texto, negrito_base=False, tam_pt=TAM):
    for tok in parse_inline_tokens(texto, negrito_base):
        if not tok['texto']: continue
        r = p.add_run(tok['texto'])
        r.bold = tok['bold']
        r.font.name = FONTE; r.font.size = Pt(tam_pt)
        if tok['yellow']:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if tok['red']:
            r.font.color.rgb = RGBColor(255, 0, 0)
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONTE)


def limpa_runs(p):
    for r in list(p._element):
        if r.tag != qn("w:pPr"): p._element.remove(r)


def paragrafo_antes(ref, texto="", negrito=False, recuo_lista=False):
    novo = copy.deepcopy(ref._element)
    ref._element.addprevious(novo)
    p = docx.text.paragraph.Paragraph(novo, ref._parent)
    limpa_runs(p)
    if recuo_lista:
        p.paragraph_format.left_indent = Pt(18); p.paragraph_format.first_line_indent = Pt(-12)
    if texto: add_runs(p, texto, negrito)
    return p


def tabela_antes(doc, ref, tab_data):
    if isinstance(tab_data, dict):
        linhas = tab_data.get("linhas", [])
        alinhamentos = tab_data.get("alinhamentos", [])
    else:
        linhas = tab_data
        alinhamentos = []
    if not linhas:
        return

    ncol = max(len(r) for r in linhas)
    t = doc.add_table(rows=len(linhas), cols=ncol)
    t.autofit = False

    secao = doc.sections[0]
    largura_conteudo = secao.page_width - secao.left_margin - secao.right_margin
    largura_twips = int(round(largura_conteudo / 635))

    tblPr = t._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(largura_twips))

    bordas = OxmlElement("w:tblBorders")
    for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{b}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
        bordas.append(e)
    tblPr.append(bordas)

    tblGrid = t._tbl.find(qn("w:tblGrid"))
    if tblGrid is None:
        tblGrid = OxmlElement("w:tblGrid")
        t._tbl.insert(t._tbl.index(tblPr) + 1, tblGrid)
    else:
        for c in list(tblGrid):
            tblGrid.remove(c)

    largura_base = largura_twips // ncol
    resto = largura_twips % ncol
    col_widths = [largura_base + (1 if j < resto else 0) for j in range(ncol)]

    for w in col_widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        tblGrid.append(col)

    for i, r in enumerate(linhas):
        trPr = t.rows[i]._tr.get_or_add_trPr()
        if i == 0:
            if trPr.find(qn("w:tblHeader")) is None:
                trPr.append(OxmlElement("w:tblHeader"))
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.append(OxmlElement("w:cantSplit"))

        for j in range(ncol):
            cel = t.cell(i, j)
            tcPr = cel._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:type"), "dxa")
            tcW.set(qn("w:w"), str(col_widths[j]))

            cel.text = ""
            cp = cel.paragraphs[0]
            val = r[j] if j < len(r) else ""

            align = alinhamentos[j] if j < len(alinhamentos) else None
            val_limpo = val.strip()
            if align is None:
                if re.match(r"^[-+]?\s*R\$\s*[\d\.,]+$", val_limpo) or re.match(r"^[-+]?[\d]{1,3}(?:\.\d{3})*,\d{2}$", val_limpo):
                    align = WD_ALIGN_PARAGRAPH.RIGHT

            if align is not None:
                cp.alignment = align

            add_runs(cp, val, negrito_base=(i == 0), tam_pt=9)
    ref._element.addprevious(t._tbl)


def imagem_antes(doc, ref, caminho_img, legenda="", largura_polegadas=6.2):
    """Insere imagem centralizada antes do parágrafo de referência com legenda opcional."""
    if not os.path.isfile(caminho_img):
        return None
    p = paragrafo_antes(ref)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(caminho_img, width=docx.shared.Inches(largura_polegadas))
    if legenda:
        p_leg = paragrafo_antes(ref)
        p_leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_leg = p_leg.add_run(legenda)
        r_leg.font.name = "Arial"
        r_leg.font.size = Pt(9.5)
        r_leg.font.italic = True
    return p


def remove(p):
    p._element.getparent().remove(p._element)


# ---------- preenchimento ----------
caminho_modelo = a.modelo
if not caminho_modelo:
    caminho_modelo = MODELO_PADRAO or os.path.join(ws_padrao(), "modelos", "MODELO RELATORIO DE INVESTIGACAO - CPJ 2026.docx")

if not os.path.exists(caminho_modelo):
    raise SystemExit(f"Modelo não encontrado: {caminho_modelo}")
doc = docx.Document(caminho_modelo)
pars = list(doc.paragraphs)
pendentes = []

for chave, rotulo in CAMPOS_CABECALHO:
    p = next((x for x in pars if x.text.strip().startswith(rotulo)), None)
    if p is None and chave == "escrivao":
        referencia = next((x for x in pars if x.text.strip().startswith("Data dos Fatos:")), None)
        if referencia is not None:
            elemento = copy.deepcopy(referencia._element)
            referencia._element.addnext(elemento)
            p = docx.text.paragraph.Paragraph(elemento, referencia._parent)
    if p is None: continue
    v = meta.get(chave, "")
    if v:
        limpa_runs(p)
        if chave in ("investigados", "vitimas"):
            v = v.upper()
        add_runs(p, f"{rotulo} {v}", negrito_base=True)
    elif chave == "escrivao":
        limpa_runs(p)
        add_runs(p, f"{rotulo} [OBTER: NOME DO ESCRIVÃO DO FEITO]", negrito_base=True)
        pendentes.append("Escrivão do feito")
    else:
        pendentes.append(rotulo.rstrip(":"))

# ---------- descoberta e preparo do fluxograma financeiro (RF23 / FD01) ----------
fluxo_png = a.fluxograma
if not fluxo_png and not a.sem_fluxograma:
    csv_candidato = a.fluxo_csv
    if not csv_candidato:
        pasta_minuta = os.path.dirname(os.path.abspath(a.minuta))
        pasta_caso = os.path.dirname(pasta_minuta)
        csv_caso = os.path.join(pasta_caso, "02-analise", "fluxo-financeiro.csv")
        if os.path.isfile(csv_caso):
            csv_candidato = csv_caso

    if csv_candidato and os.path.isfile(csv_candidato):
        try:
            import importlib.util
            script_diag = os.path.join(os.path.dirname(__file__), "gerar_diagrama_financeiro.py")
            if os.path.isfile(script_diag):
                spec = importlib.util.spec_from_file_location("gerar_diagrama_financeiro", script_diag)
                mod_diag = importlib.util.module_from_spec(spec)
                spec.loader.exec_module(mod_diag)

                pasta_caso_nome = os.path.basename(pasta_caso)
                if pasta_caso_nome.startswith("OS-"):
                    caso_id = pasta_caso_nome
                elif meta.get("ordem_servico"):
                    clean = re.sub(r"[^\w]+", "-", str(meta["ordem_servico"])).strip("-")
                    caso_id = f"OS-{clean}" if clean else pasta_caso_nome
                else:
                    caso_id = pasta_caso_nome

                saida_diag = os.path.join(os.path.dirname(os.path.abspath(a.saida)), f"FLUXO-FINANCEIRO-{caso_id}.png")
                txs = mod_diag.carregar_transacoes(csv_candidato)
                if txs:
                    fluxo_png = mod_diag.desenhar_diagrama(
                        txs,
                        saida_diag,
                        titulo="FLUXOGRAMA DO CAMINHO DO DINHEIRO — RASTREABILIDADE BANCÁRIA",
                        caso_id=caso_id,
                        dpi=300,
                    )
        except Exception as e:
            print(f"Aviso ao gerar fluxograma: {e}", file=sys.stderr)

fluxo_injetado = False

for secao, prefixos in SECOES.items():
    guias = [x for x in pars if any(x.text.strip().startswith(pr) for pr in prefixos)]
    if not guias: continue
    conteudo = blocos(secoes.get(secao, []))
    if not conteudo:
        pendentes.append(secao); continue
    ref = guias[0]
    for n, (tipo, val) in enumerate(conteudo):
        if n: paragrafo_antes(ref)  # linha em branco entre blocos, como no modelo
        if tipo == "tab": tabela_antes(doc, ref, val)
        elif tipo == "h": paragrafo_antes(ref, val, negrito=True)
        elif tipo == "li": paragrafo_antes(ref, val, recuo_lista=True)
        elif tipo == "img":
            legenda, arq_img = val
            if not os.path.isabs(arq_img):
                arq_img = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(a.minuta)), arq_img))
            imagem_antes(doc, ref, arq_img, legenda=legenda)
            if fluxo_png and os.path.abspath(arq_img) == os.path.abspath(fluxo_png):
                fluxo_injetado = True
        else: paragrafo_antes(ref, val)

    # Injeção automática na seção DILIGÊNCIAS REALIZADAS se houver fluxo e não estiver explícito
    if (secao == "DILIGÊNCIAS REALIZADAS" or (secao == "CONCLUSÃO" and not fluxo_injetado)) and fluxo_png and not fluxo_injetado:
        paragrafo_antes(ref)
        paragrafo_antes(ref, "Fluxograma do Caminho do Dinheiro e Repasses Bancários", negrito=True)
        imagem_antes(
            doc,
            ref,
            fluxo_png,
            legenda="Figura 1 — Fluxograma do Caminho do Dinheiro e Rastreabilidade das Camadas Bancárias (300 DPI)",
        )
        paragrafo_antes(ref)
        fluxo_injetado = True

    for g in guias: remove(g)

for chave, prefixo in (("local_data", "[local, Estado]"), ("delegado", "{Nome do Delegado")):
    p = next((x for x in doc.paragraphs if x.text.strip().startswith(prefixo)), None)
    if p is None: continue
    if meta.get(chave):
        limpa_runs(p); add_runs(p, meta[chave], negrito_base=(chave == "delegado"))
    else:
        pendentes.append(chave)

# Tratamento do(a) delegado(a): concordância de gênero (saudação inicial e endereçamento final).
# Campo da minuta: delegado_genero: M | F (masculino/feminino). Fonte: documentos do caso (Dr./Dra., "o/a Delegado/a") ou o operador.
gen = (meta.get("delegado_genero") or "").strip().lower()[:1]
if gen in ("m", "f"):
    fem = gen == "f"
    troca = {
        "EXCELENTÍSSIMO (A) SENHOR (A)": "EXCELENTÍSSIMA SENHORA DOUTORA DELEGADA DE POLÍCIA," if fem else "EXCELENTÍSSIMO SENHOR DOUTOR DELEGADO DE POLÍCIA,",
        "A(o) Excelentíssimo (a)": "À Excelentíssima Sra. Dra." if fem else "Ao Excelentíssimo Sr. Dr.",
        "Delegado (a) de Polícia Civil": "Delegada de Polícia" if fem else "Delegado de Polícia",
    }
    for prefixo, novo in troca.items():
        p = next((x for x in doc.paragraphs if x.text.strip().startswith(prefixo)), None)
        if p is not None:
            limpa_runs(p); add_runs(p, novo, negrito_base=True)
else:
    pendentes.append("delegado_genero (DEFINIR SE O DELEGADO É HOMEM OU MULHER: M ou F — ajusta saudação e endereçamento)")

if meta.get("data_rodape"):
    for sec in doc.sections:
        for par in [q for t in sec.footer.tables for c in t._cells for q in c.paragraphs] + list(sec.footer.paragraphs):
            for r in par.runs:
                if re.fullmatch(r"\s*\d{2}/\d{2}/\d{4}\s*", r.text or ""):
                    r.text = meta["data_rodape"]

for i, p in enumerate(doc.paragraphs):
    if p._element.xpath(".//pic:pic") and not p.text.strip():
        if a.sem_assinatura:
            if i + 1 < len(doc.paragraphs):
                doc.paragraphs[i + 1].paragraph_format.page_break_before = True
            remove(p)
        else:
            p.paragraph_format.page_break_before = True

# Conferência final: texto-guia do modelo, chaves e marcadores que não podem chegar ao delegado.
texto_doc = "\n".join([p.text for p in doc.paragraphs] +
                      [q.text for t in doc.tables for c in t._cells for q in c.paragraphs])
residuos = sorted({m.group(0) for m in re.finditer(r" \((?:A|a)\)|A\(o\)|\{[^}\n]{0,40}\}?|\[n[º°o][^\]\n]{0,20}\]?|\bPREENCHER\b|\bTODO\b", texto_doc)})
if residuos:
    pendentes.append("RESÍDUOS NO TEXTO (CORRIGIR ANTES DE ENTREGAR): " + " | ".join(residuos[:10]))

os.makedirs(os.path.dirname(os.path.abspath(a.saida)), exist_ok=True)
doc.save(a.saida)
print(f"OK -> {a.saida}")
if pendentes:
    print("Campos mantidos com texto-guia do modelo (preencher):", ", ".join(pendentes))
