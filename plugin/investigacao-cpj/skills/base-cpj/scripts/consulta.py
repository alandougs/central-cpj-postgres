#!/usr/bin/env python3
"""Bases de CONSULTA (ex.: exportações do Muralha Paulista em Excel/Word/CSV) — somente para pesquisa.
NUNCA são fonte do relatório de investigação: o relatório vem apenas do IP/peças enviadas ao caso.

Uso (CLI):
  consulta.py importar ARQUIVO [--nome "Muralha Paulista - set/2026"]
  consulta.py listar
  consulta.py remover BASE_ID

Guarda em consulta\\<BASE_ID>\\: o arquivo original, registros.jsonl (1 registro por linha) e base.json
(metadados e mapeamento de colunas). O indexar.py carrega os registros na tabela `pessoas` do RAG.
Colunas reconhecidas por sinônimo: nome, mãe, pai, CPF, RG, nascimento, telefones, endereços, empresas,
CNPJ, e-mails, placas. Colunas não reconhecidas são preservadas em `extras` e entram na busca textual.
"""
import argparse, csv, datetime, json, os, re, shutil, unicodedata

WS = os.environ.get("CPJ_WORKSPACE", r"C:\CPJ - TRABALHO")
BASES = os.path.join(WS, "consulta")
CAMPOS = ["nome", "mae", "pai", "cpf", "rg", "nascimento", "telefones", "enderecos", "empresas", "cnpj", "emails", "placas",
          "veiculos", "processos", "bos", "mandados", "cautelares"]
SEP = {"enderecos": ", "}  # separador ao juntar valores; demais: " | "
SINONIMOS = {
    "nome": ["nome", "nome completo", "nome civil", "nome da pessoa", "pessoa", "individuo", "nome do individuo",
             "envolvido", "nome social", "investigado", "vitima", "titular", "qualificado"],
    "mae": ["mae", "nome da mae", "nome mae", "genitora", "filiacao mae", "filiacao materna"],
    "pai": ["pai", "nome do pai", "nome pai", "genitor", "filiacao pai", "filiacao paterna"],
    "cpf": ["cpf", "n cpf", "no cpf", "numero cpf", "numero do cpf", "cpf do titular"],
    "rg": ["rg", "registro geral", "identidade", "n rg", "no rg", "numero rg", "rg/uf"],
    "nascimento": ["nascimento", "data de nascimento", "data nascimento", "dt nasc", "dt nascimento", "data nasc", "nasc"],
    "telefones": ["telefone", "telefones", "celular", "celulares", "fone", "fones", "tel", "contato", "whatsapp", "telefone celular"],
    "enderecos": ["endereco", "enderecos", "logradouro", "residencia", "rua", "bairro", "cidade", "municipio", "uf", "cep",
                  "complemento", "endereco residencial", "endereco completo"],
    "empresas": ["empresa", "empresas", "razao social", "nome fantasia", "empregador", "sociedade"],
    "cnpj": ["cnpj", "n cnpj", "numero cnpj"],
    "emails": ["email", "e-mail", "e mail", "correio eletronico"],
    "placas": ["placa", "placas", "placa do veiculo", "veiculo placa"],
    "veiculos": ["veiculo", "veiculos", "marca modelo", "marca/modelo", "modelo do veiculo"],
    "processos": ["processo", "processos", "inquerito", "inqueritos", "inquerito policial", "ip", "tc", "termo circunstanciado",
                  "procedimento", "procedimentos", "antecedentes", "antecedentes criminais", "autos", "vpi"],
    "bos": ["bo", "bos", "boletim", "boletins", "boletim de ocorrencia", "boletins de ocorrencia", "rdo", "ocorrencias"],
    "mandados": ["mandado", "mandados", "mandado de prisao", "mandados de prisao", "situacao mandado"],
    "cautelares": ["medida cautelar", "medidas cautelares", "cautelar", "cautelares", "medida protetiva", "medidas protetivas"],
}

# ---------------------------------------------------------------- antecedentes e identificadores em texto livre
RX = {
    "cpf": re.compile(r"\b\d{3}\.\d{3}\.\d{3}-\d{2}\b"),
    "telefones": re.compile(r"\(?\b\d{2}\)?\s?9?\d{4}-\d{4}\b"),
    "placas": re.compile(r"\b[A-Z]{3}-?\d[A-Z0-9]\d{2}\b"),
    "cnpj": re.compile(r"\b\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}\b"),
    "cnj": re.compile(r"\b\d{7}-\d{2}\.\d{4}\.\d\.\d{2}\.\d{4}\b"),
}
CAUTELAR = ("medida cautelar", "medidas cautelares", "medida protetiva", "medidas protetivas", "monitoramento eletronico",
            "tornozeleira", "recolhimento domiciliar", "proibicao de contato", "comparecimento periodico", "comparecimento mensal",
            "liberdade provisoria", "fianca", "prisao domiciliar")


# ---------------------------------------------------------------- estado de mandados e cautelares (L01)
# O trecho original é sempre preservado; o estado é uma INTERPRETAÇÃO a conferir, nunca prova de vigência.
#   confirmado    = o texto afirma situação atual (em aberto, vigente, a cumprir, não cumprido, foragido...)
#   historico     = o texto indica encerramento (cumprido, revogado, recolhido, extinto, sem efeito...)
#   negado        = o texto nega a existência (sem mandado, nada consta, "Não" em coluna própria...)
#   indeterminado = menção sem situação declarada (ex.: "Mandado de prisão nº 123", "Sim", "expedido")
ESTADOS = ("confirmado", "indeterminado", "historico", "negado")
NEG_VALOR = {"nao", "n", "nenhum", "nenhuma", "nada", "nada consta", "negativo", "negativa", "inexistente", "inexiste",
             "sem registro", "sem registros", "sem", "0", "nao consta", "nao possui", "nao ha", "ausente", "nao existe"}
RX_NEG = re.compile(r"\b(sem|nenhum|nenhuma|nao ha|nao consta|nao constam|nao possui|nao existe|nao existem|nao foi localizad\w*|"
                    r"nao foram localizad\w*|inexist\w*|ausencia de|livre de)\b(\s+\w+){0,4}?\s+"
                    r"(mandado|mandados|medida|medidas|cautelar\w*|protetiv\w*|restric\w*|registro\w*|pendencia\w*)")
RX_NADA = re.compile(r"\b(nada consta|negativ[oa]s?|sem registros?|sem restric\w*|sem pendencias?)\b")
RX_HIST = re.compile(r"\b(cumprid[oa]s?|recolhid[oa]s?|revogad[oa]s?|baixad[oa]s?|extint[oa]s?|cessad[oa]s?|expirad[oa]s?|"
                     r"arquivad[oa]s?|contramandad[oa]s?|contramandado|prescrit[oa]s?|cancelad[oa]s?|sem efeito|encerrad[oa]s?|"
                     r"finalizad[oa]s?|cumprimento encerrado)\b")
RX_ATUAL = re.compile(r"\b(em aberto|aberto|aberta|pendente de cumprimento|a cumprir|vigente\w*|em vigor|ativ[oa]s?|"
                      r"foragid[oa]|procurad[oa]|em cumprimento)\b")
RX_NAO_ENCERRADO = re.compile(r"\bnao (cumprid|recolhid|revogad|baixad|extint)\w*")


def estado_antecedente(trecho):
    """Classifica UMA menção (linha/célula) de mandado ou cautelar. Ver ESTADOS."""
    n = norm(trecho)
    if not n: return "indeterminado"
    if ":" in str(trecho):
        val = norm(str(trecho).split(":", 1)[1])
        if val in NEG_VALOR: return "negado"
    if RX_NAO_ENCERRADO.search(n): return "confirmado"
    hist = bool(RX_HIST.search(n))
    if RX_NADA.search(n) or RX_NEG.search(n):
        return "historico" if hist else "negado"   # "sem mandado em aberto; anterior cumprido" -> histórico
    if hist: return "historico"
    if RX_ATUAL.search(n): return "confirmado"
    return "indeterminado"


def estado_geral(estados):
    """Estado do registro: prevalece o mais relevante para a investigação."""
    for e in ESTADOS:
        if e in estados: return e
    return ""


def _itens(reg):
    """Reconstrói os itens de antecedentes a partir dos trechos preservados nos campos mandados/cautelares."""
    itens = []
    for campo, tipo in (("mandados", "mandado"), ("cautelares", "cautelar")):
        ests = []
        for t in [x.strip() for x in re.split(r"\s\|\s", reg.get(campo) or "") if x.strip()]:
            e = estado_antecedente(t); ests.append(e)
            itens.append({"tipo": tipo, "trecho": t, "estado": e, "localizador": reg.get("localizador", "")})
        reg[f"{campo}_estado"] = estado_geral(ests)
    reg["antecedentes_itens"] = itens
    return reg


CAMPOS_COM_CONTEXTO = ("mandados", "cautelares", "processos", "bos")  # guardam a linha inteira, não só o valor


def _add(reg, campo, valor):
    valor = (valor or "").strip()
    if not valor: return
    atuais = [x.strip() for x in re.split(r"\s\|\s", reg.get(campo) or "") if x.strip()]
    if norm(valor) not in {norm(x) for x in atuais}:
        reg[campo] = SEP.get(campo, " | ").join(atuais + [valor[:220]])


def enriquecer(reg, texto):
    """Extrai de um bloco de texto os antecedentes (processos/IP/TC, BOs, mandados, cautelares) linha a linha.
    Identificadores (CPF, telefone, placa, CNPJ) só preenchem o campo se ele estiver VAZIO — evita misturar terceiros."""
    for ln in (texto or "").splitlines():
        n = norm(ln); s = ln.strip()
        if not n: continue
        if "mandado" in n and "prisao" in n: _add(reg, "mandados", s)
        if any(c in n for c in CAUTELAR): _add(reg, "cautelares", s)
        tem_num = bool(re.search(r"\d[\d./-]{3,}", ln))
        if RX["cnj"].search(ln) or (tem_num and re.search(r"\b(inquerito|ip|tc|termo circunstanciado|processo|autos|vpi|execucao)\b", n)):
            _add(reg, "processos", s)
        if tem_num and re.search(r"\b(bo|b o|boletim de ocorrencia|boletim|rdo)\b", n): _add(reg, "bos", s)
    for campo in ("cpf", "telefones", "placas", "cnpj"):
        if not (reg.get(campo) or "").strip():
            for m in dict.fromkeys(RX[campo].findall(texto or "")): _add(reg, campo, m)
    return _itens(reg)


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9@/ -]+", " ", s).replace("-", " ").strip()


def campo_de(cabecalho):
    h = re.sub(r"\s+", " ", norm(cabecalho))
    if not h: return None
    for campo, sins in SINONIMOS.items():
        if h in sins: return campo
    for campo, sins in SINONIMOS.items():
        if any(h.startswith(s + " ") or h.endswith(" " + s) for s in sins if len(s) > 3): return campo
    return None


def slug(s): return re.sub(r"[^a-z0-9]+", "-", norm(s)).strip("-")[:40] or "base"


def celula(v):
    if v is None: return ""
    if isinstance(v, (datetime.date, datetime.datetime)): return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------- leitores (geram listas de linhas)
def tabelas_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        abas = [(ws.title, [[celula(c) for c in r] for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()  # read_only mantém o arquivo aberto no Windows até fechar
    yield from abas


def tabelas_xls(p):
    import xlrd
    wb = xlrd.open_workbook(p)
    for sh in wb.sheets():
        yield sh.name, [[celula(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]


def tabelas_csv(p):
    raw = open(p, encoding="utf-8-sig", errors="replace").read()
    try: dia = csv.Sniffer().sniff(raw[:5000], delimiters=";,\t|")
    except csv.Error: dia = csv.excel; dia.delimiter = ";"
    yield "csv", [[c.strip() for c in r] for r in csv.reader(raw.splitlines(), dia)]


def docx_conteudo(p):
    import docx
    d = docx.Document(p)
    tabs = [("tabela %d" % (i + 1), [[c.text.strip() for c in r.cells] for r in t.rows]) for i, t in enumerate(d.tables)]
    return tabs, [par.text for par in d.paragraphs]


def registros_de_tabela(nome_aba, linhas):
    """Detecta a linha de cabeçalho (até 20 primeiras) e converte as demais em registros."""
    ini = None
    for i, r in enumerate(linhas[:20]):
        if sum(1 for c in r if campo_de(c)) >= 1 and sum(1 for c in r if c) >= 2: ini = i; break
    if ini is None: return [], {}
    cab = linhas[ini]
    mapa = {j: campo_de(h) for j, h in enumerate(cab)}
    out = []
    for k, r in enumerate(linhas[ini + 1:], start=ini + 2):
        if not any(c for c in r): continue
        reg = {c: "" for c in CAMPOS}; extras = {}
        for j, v in enumerate(r):
            if not v: continue
            h = cab[j] if j < len(cab) and cab[j] else f"coluna {j + 1}"
            if mapa.get(j) in CAMPOS_COM_CONTEXTO: _add(reg, mapa[j], f"{h}: {v}")   # "Mandado de prisão: Não" ≠ "Não"
            elif mapa.get(j): _add(reg, mapa[j], v)
            else: extras[h] = v
        texto = "\n".join(f"{cab[j] if j < len(cab) else ''}: {v}" for j, v in enumerate(r) if v)
        enriquecer(reg, texto)
        out.append({"localizador": f"{nome_aba}, linha {k}", **reg, "extras": extras, "texto": texto[:20000]})
    return out, {cab[j]: m for j, m in mapa.items() if m and j < len(cab)}


def registros_de_texto(linhas, rotulo="linha"):
    """Texto livre (Word sem tabela, PDF, TXT, texto colado). Linhas 'Campo: valor' preenchem campos; um novo
    registro começa quando reaparece o campo Nome. Todo o bloco de cada registro é guardado e varrido para
    antecedentes (processos/IP/TC, BOs, mandados de prisão, medidas cautelares). Sem nenhum 'Nome:', o texto
    inteiro vira um único registro pesquisável."""
    out, atual = [], None
    def novo(i): return {**{c: "" for c in CAMPOS}, "localizador": f"{rotulo} {i}", "extras": {}, "_linhas": []}
    for i, ln in enumerate(linhas, 1):
        m = re.match(r"^\s*([^:]{2,40}):\s*(.+)$", ln or "")
        campo = campo_de(m.group(1)) if m else None
        if campo == "nome" and atual and atual.get("nome"): out.append(atual); atual = None
        if atual is None: atual = novo(i)
        atual["_linhas"].append(ln or "")
        if m:
            if campo in CAMPOS_COM_CONTEXTO: _add(atual, campo, ln.strip())
            elif campo: _add(atual, campo, m.group(2))
            else: atual["extras"][m.group(1).strip()] = m.group(2).strip()
    if atual: out.append(atual)
    regs = []
    for r in out:
        texto = "\n".join(r.pop("_linhas")).strip()
        if not texto: continue
        r["texto"] = texto[:20000]; enriquecer(r, texto); regs.append(r)
    return regs


def linhas_pdf(p):
    """Texto do PDF por página; páginas sem texto passam por OCR (Tesseract português), se disponível."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(p); out = []
    for i in range(len(pdf)):
        t = pdf[i].get_textpage().get_text_range().strip()
        if len(t) < 30:
            try:
                import pytesseract
                tess = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
                if os.path.exists(tess): pytesseract.pytesseract.tesseract_cmd = tess
                td = os.path.join(WS, "ferramentas", "tessdata")
                cfg = f'--tessdata-dir "{td}"' if os.path.exists(os.path.join(td, "por.traineddata")) else ""
                t = pytesseract.image_to_string(pdf[i].render(scale=200 / 72).to_pil(), lang="por", config=cfg)
            except Exception:
                t = ""
        out += t.splitlines()
    return out


# ---------------------------------------------------------------- operações
FORMATOS = (".xlsx", ".xlsm", ".xls", ".csv", ".docx", ".pdf", ".txt")


def importar(arquivo, nome=None, progresso=None):
    ext = os.path.splitext(arquivo)[1].lower()
    if ext not in FORMATOS:
        raise ValueError("Formato não suportado (use Excel, CSV, Word .docx, PDF ou TXT). Arquivo .doc antigo: salve como .docx.")
    nome = nome or os.path.splitext(os.path.basename(arquivo))[0]
    base_id = f"{datetime.date.today():%Y%m%d}-{slug(nome)}"
    k = 2
    while os.path.exists(os.path.join(BASES, base_id)): base_id = f"{base_id.rsplit('_', 1)[0]}_{k}"; k += 1
    destino = os.path.join(BASES, base_id); os.makedirs(destino)
    if progresso: progresso(10, "lendo arquivo")
    regs, mapeamento, fontes = [], {}, []
    try:
        if ext in (".xlsx", ".xlsm"): fontes = list(tabelas_xlsx(arquivo))
        elif ext == ".xls": fontes = list(tabelas_xls(arquivo))
        elif ext == ".csv": fontes = list(tabelas_csv(arquivo))
        elif ext == ".docx":
            tabs, pars = docx_conteudo(arquivo); fontes = tabs
            regs += registros_de_texto(pars, "parágrafo")
        elif ext == ".pdf":
            if progresso: progresso(20, "lendo PDF (OCR se necessário)")
            regs += registros_de_texto(linhas_pdf(arquivo))
        else:
            regs += registros_de_texto(open(arquivo, encoding="utf-8-sig", errors="replace").read().splitlines())
        for i, (aba, linhas) in enumerate(fontes):
            r, m = registros_de_tabela(aba, linhas); regs += r; mapeamento.update(m)
            if progresso: progresso(10 + int(70 * (i + 1) / max(1, len(fontes))), f"interpretando {aba}")
    except Exception:
        shutil.rmtree(destino, ignore_errors=True); raise
    if not regs:
        shutil.rmtree(destino)
        raise ValueError("Nenhum registro reconhecido. Em planilhas, use cabeçalho com colunas como Nome, Mãe, CPF, Telefone; "
                         "em texto, linhas no formato 'Campo: valor'.")
    shutil.copyfile(arquivo, os.path.join(destino, "original" + ext))
    with open(os.path.join(destino, "registros.jsonl"), "w", encoding="utf-8") as f:
        for r in regs: f.write(json.dumps(r, ensure_ascii=False) + "\n")
    resumo = {c: sum(1 for r in regs if (r.get(c) or "").strip()) for c in ("processos", "bos")}
    for c in ("mandados", "cautelares"):  # por estado interpretado (a conferir)
        resumo[c] = {e: sum(1 for r in regs if r.get(f"{c}_estado") == e) for e in ESTADOS}
    meta = {"id": base_id, "nome": nome, "arquivo_original": os.path.basename(arquivo), "importado_em":
            datetime.datetime.now().isoformat(timespec="seconds"), "registros": len(regs), "mapeamento": mapeamento,
            "com_antecedentes": resumo, "uso": "somente consulta — não é fonte do relatório de investigação"}
    json.dump(meta, open(os.path.join(destino, "base.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    if progresso: progresso(90, "concluído")
    return meta


def importar_texto(texto, nome, progresso=None):
    """Texto colado (ex.: ficha copiada de um sistema) — vira uma base de consulta."""
    if not (texto or "").strip(): raise ValueError("Cole o texto a importar.")
    tmp = os.path.join(BASES, f"_colado-{datetime.datetime.now():%Y%m%d%H%M%S%f}.txt"); os.makedirs(BASES, exist_ok=True)
    open(tmp, "w", encoding="utf-8").write(texto)
    try: return importar(tmp, nome or "Texto colado", progresso)
    finally:
        try: os.remove(tmp)
        except OSError: pass


def listar():
    if not os.path.isdir(BASES): return []
    out = []
    for d in sorted(os.listdir(BASES)):
        p = os.path.join(BASES, d, "base.json")
        if os.path.exists(p): out.append(json.load(open(p, encoding="utf-8")))
    return out


def remover(base_id):
    d = os.path.normpath(os.path.join(BASES, base_id))
    if not d.startswith(os.path.normpath(BASES) + os.sep) or not os.path.isdir(d): raise FileNotFoundError(base_id)
    shutil.rmtree(d)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("importar"); s.add_argument("arquivo"); s.add_argument("--nome")
    sub.add_parser("listar")
    s = sub.add_parser("remover"); s.add_argument("base_id")
    a = ap.parse_args()
    if a.cmd == "importar":
        m = importar(a.arquivo, a.nome); print(f"Base {m['id']}: {m['registros']} registros. Mapeamento: {m['mapeamento']}")
        print("Rode indexar.py para torná-la pesquisável.")
    elif a.cmd == "listar":
        for m in listar(): print(f"{m['id']:<40} {m['registros']:>7} registros  ({m['arquivo_original']})")
    else:
        remover(a.base_id); print("Base removida. Rode indexar.py.")