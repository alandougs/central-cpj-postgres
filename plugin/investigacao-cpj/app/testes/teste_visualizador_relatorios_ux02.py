#!/usr/bin/env python3
"""Prévia DOCX e pasta de relatórios em workspace fictício; sem servidor ou IA."""
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import quote
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE


class VisualizadorRelatoriosUX02(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="cpj-previa-ux02-")
        cls.ws = Path(cls.temp.name)
        cls.env = {k: os.environ.get(k) for k in ("CPJ_WORKSPACE", "CPJ_SEM_AGENTE_EMBUTIDO")}
        os.environ["CPJ_WORKSPACE"] = str(cls.ws)
        os.environ["CPJ_SEM_AGENTE_EMBUTIDO"] = "1"
        (cls.ws / "casos").mkdir()
        sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
        import servidor
        import rotas.casos as casos
        cls.s = servidor
        cls.rotas = casos
        assert Path(servidor.WS).resolve() == cls.ws.resolve(), "Execute em processo isolado."
        cls.s.app.config.update(TESTING=True)
        cls.id = "OS-901-2099"
        cls.pasta = cls.ws / "casos" / cls.id
        cls.relatorios = cls.pasta / "03-relatorios"
        cls.relatorios.mkdir(parents=True)
        cls.originais = cls.pasta / "00-originais"
        cls.originais.mkdir()
        cls.nome = "Relatório 'fictício' & teste.docx"
        doc = Document()
        p = doc.add_paragraph("PARAGRAFO_INICIAL ")
        p.add_run("<script>alert('teste')</script> & texto").bold = True
        p.add_run(" ITALICO_FICTICIO").italic = True
        tabela = doc.add_table(rows=1, cols=2)
        tabela.cell(0, 0).text = "TABELA_INTERMEDIARIA <img src=x onerror=teste>"
        tabela.cell(0, 1).text = "Célula fictícia"
        doc.add_paragraph("PARAGRAFO_FINAL")
        p = doc.add_paragraph()
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), doc.part.relate_to("https://externo.invalid/teste", RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        texto = OxmlElement("w:t")
        texto.text = "LINK_FICTICIO"
        run.append(texto)
        link.append(run)
        p._p.append(link)
        doc.save(cls.relatorios / cls.nome)
        doc.save(cls.relatorios / "RELATORIO-FINAL.docx")
        doc.save(cls.originais / "original.docx")
        cls.pdf = cls.relatorios / "RELATORIO-FINAL.pdf"
        cls.pdf.write_bytes(b"%PDF-1.4\n% ficticio\n%%EOF")
        (cls.relatorios / "minuta.pdf").write_bytes(cls.pdf.read_bytes())
        (cls.relatorios / "texto.txt").write_text("Fictício", encoding="utf-8")
        cls.irma = cls.ws / "casos" / (cls.id + "-irma")
        (cls.irma / "03-relatorios").mkdir(parents=True)
        doc.save(cls.irma / "03-relatorios" / "RELATORIO-FINAL.docx")

    @classmethod
    def tearDownClass(cls):
        for nome, valor in cls.env.items():
            if valor is None:
                os.environ.pop(nome, None)
            else:
                os.environ[nome] = valor
        cls.temp.cleanup()

    def setUp(self):
        self.cliente = self.s.app.test_client()
        self.usuario = {"login": "investigador_ficticio", "perfil": "investigador"}
        self.auth = patch("rotas.comum.usuario", side_effect=lambda: self.usuario)
        self.auth.start()
        self.addCleanup(self.auth.stop)

    def previa(self, rel=None):
        return self.cliente.get("/visualizar/" + self.id + "/" + quote(rel or ("03-relatorios/" + self.nome), safe="/"))

    def abrir(self, sub="03-relatorios", remoto="127.0.0.1"):
        return self.cliente.post("/api/casos/" + self.id + "/abrir", json={"sub": sub},
                                 headers={"X-CPJ": "1"}, environ_overrides={"REMOTE_ADDR": remoto})

    def test_docx_preserva_ordem_texto_tabela_e_formatacao_simples(self):
        r = self.previa()
        self.assertEqual(r.status_code, 200)
        texto = r.get_data(as_text=True)
        self.assertLess(texto.index("PARAGRAFO_INICIAL"), texto.index("TABELA_INTERMEDIARIA"))
        self.assertLess(texto.index("TABELA_INTERMEDIARIA"), texto.index("PARAGRAFO_FINAL"))
        self.assertIn("<table><tbody>", texto)
        self.assertIn("<strong>&lt;script&gt;", texto)
        self.assertIn("<em> ITALICO_FICTICIO</em>", texto)
        self.assertIn("LINK_FICTICIO", texto)
        self.assertNotIn("https://externo.invalid", texto)
        self.assertIn("Prévia de texto e tabelas", texto)
        self.assertIn("A formatação completa está no arquivo original", texto)
        self.assertIn("frame-ancestors 'self'", r.headers["Content-Security-Policy"])
        self.assertEqual(r.headers["Cache-Control"], "no-store")

    def test_escapa_texto_nome_e_link_do_original(self):
        r = self.previa()
        texto = r.get_data(as_text=True)
        self.assertIn("Relatório &#x27;fictício&#x27; &amp; teste.docx", texto)
        self.assertIn("&lt;img src=x onerror=teste&gt;", texto)
        self.assertNotIn("<script>", texto)
        self.assertNotIn("<img ", texto)
        self.assertIn("href='/arquivo/", texto)
        self.assertNotIn("href='/arquivo/" + self.id + "/03-relatorios/Relatório 'fictício'", texto)

    def test_original_word_baixa_e_pdf_serve_no_navegador(self):
        word = self.cliente.get("/arquivo/" + self.id + "/03-relatorios/" + quote(self.nome))
        with word:
            self.assertEqual(word.status_code, 200)
            self.assertIn("attachment", word.headers["Content-Disposition"])
            self.assertTrue(word.get_data().startswith(b"PK"))
        pdf = self.cliente.get("/arquivo/" + self.id + "/03-relatorios/RELATORIO-FINAL.pdf")
        with pdf:
            self.assertEqual(pdf.status_code, 200)
            self.assertEqual(pdf.mimetype, "application/pdf")
            self.assertNotIn("attachment", pdf.headers["Content-Disposition"])
            self.assertEqual(pdf.get_data(), self.pdf.read_bytes())

    def test_minutas_restritas_e_finais_permitidos_para_gestao(self):
        for perfil in ("delegado", "escrivao"):
            with self.subTest(perfil=perfil):
                self.usuario["perfil"] = perfil
                self.assertEqual(self.previa().status_code, 403)
                self.assertEqual(self.previa("03-relatorios/RELATORIO-FINAL.docx").status_code, 200)
                negado = self.cliente.get("/arquivo/" + self.id + "/03-relatorios/minuta.pdf")
                self.assertEqual(negado.status_code, 403)
                final = self.cliente.get("/arquivo/" + self.id + "/03-relatorios/RELATORIO-FINAL.pdf")
                with final:
                    self.assertEqual(final.status_code, 200)

    def test_exige_sessao_e_permissao(self):
        self.usuario = None
        self.assertEqual(self.previa().status_code, 401)
        self.usuario = {"login": "sem_permissao", "perfil": "inexistente"}
        self.assertEqual(self.previa().status_code, 403)

    def test_rejeita_extensao_e_documento_fora_de_relatorios(self):
        for rel in ("03-relatorios/RELATORIO-FINAL.pdf", "03-relatorios/texto.txt", "00-originais/original.docx",
                    "03-relatorios/../00-originais/original.docx"):
            with self.subTest(rel=rel):
                self.assertEqual(self.previa(rel).status_code, 404)

    def test_rejeita_caminho_externo_prefixo_irmao_e_id_parental(self):
        rel = "../" + self.irma.name + "/03-relatorios/RELATORIO-FINAL.docx"
        self.assertEqual(self.previa(rel).status_code, 404)
        arquivo = self.cliente.get("/arquivo/" + self.id + "/" + quote(rel, safe="/"))
        self.assertEqual(arquivo.status_code, 404)
        rel_windows = "..\\" + self.irma.name + "\\03-relatorios\\RELATORIO-FINAL.docx"
        self.assertEqual(self.previa(rel_windows).status_code, 404)
        self.assertEqual(self.cliente.get("/visualizar/../03-relatorios/RELATORIO-FINAL.docx").status_code, 404)

    def test_abrir_pasta_de_relatorios_e_bloquear_escape(self):
        with patch.object(self.rotas.os, "startfile", create=True) as abrir:
            self.assertEqual(self.abrir().status_code, 200)
            abrir.assert_called_once_with(os.path.realpath(self.relatorios))
            abrir.reset_mock()
            for sub in ("../" + self.irma.name, str(self.irma), "..", "03-relatorios/inexistente", 42):
                with self.subTest(sub=sub):
                    self.assertEqual(self.abrir(sub).status_code, 404)
            abrir.assert_not_called()

    def test_abrir_restrito_a_local_trabalho_e_anti_csrf(self):
        with patch.object(self.rotas.os, "startfile", create=True) as abrir:
            self.assertEqual(self.abrir(remoto="192.0.2.10").status_code, 403)
            self.usuario["perfil"] = "delegado"
            self.assertEqual(self.abrir().status_code, 403)
            self.usuario["perfil"] = "investigador"
            r = self.cliente.post("/api/casos/" + self.id + "/abrir", json={"sub": "03-relatorios"})
            self.assertEqual(r.status_code, 403)
            abrir.assert_not_called()

    def test_limites_de_tamanho_expansao_xml_e_html(self):
        for limite in ("PREVIA_DOCX_BYTES", "PREVIA_DOCX_EXPANDIDO", "PREVIA_DOCX_XML", "PREVIA_DOCX_HTML"):
            with self.subTest(limite=limite), patch.object(self.rotas, limite, 1):
                self.assertEqual(self.previa().status_code, 413)

    def test_rejeita_docx_corrompido_e_compressao_excessiva(self):
        (self.relatorios / "corrompido.docx").write_bytes(b"DOCX ficticio invalido")
        self.assertEqual(self.previa("03-relatorios/corrompido.docx").status_code, 422)
        with zipfile.ZipFile(self.relatorios / "comprimido.docx", "w", compression=zipfile.ZIP_DEFLATED) as pacote:
            pacote.writestr("word/document.xml", "x" * (2 * 1024 * 1024))
        self.assertEqual(self.previa("03-relatorios/comprimido.docx").status_code, 413)

    def test_realpath_rejeita_atalho_para_fora_do_caso(self):
        alvo = self.relatorios / "RELATORIO-FINAL.docx"
        original = self.rotas.os.path.realpath
        with patch.object(self.rotas.os.path, "realpath", side_effect=lambda p: original(self.irma / "03-relatorios" / alvo.name)
                          if os.path.normcase(original(p)) == os.path.normcase(original(alvo)) else original(p)):
            self.assertEqual(self.previa("03-relatorios/RELATORIO-FINAL.docx").status_code, 404)


if __name__ == "__main__":
    unittest.main(verbosity=2)
