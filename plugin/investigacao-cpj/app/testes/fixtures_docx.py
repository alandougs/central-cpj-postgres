"""Modelo sintético para sandboxes de teste; nunca usado como fallback em produção."""
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw

NOME = "MODELO RELATORIO DE INVESTIGACAO - CPJ 2026.docx"


def criar_modelo(pasta):
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    destino = pasta / NOME
    if destino.exists():
        return destino
    doc = Document()
    doc.core_properties.title = "FIXTURE SINTÉTICA CPJ — SOMENTE TESTES"
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    img = Image.new("RGB", (200, 40), "white")
    ImageDraw.Draw(img).rectangle((5, 5, 195, 35), outline="black")
    png = BytesIO()
    img.save(png, format="PNG")
    png.seek(0)
    h = sec.header.add_table(rows=1, cols=2, width=Cm(17))
    h.cell(0, 0).paragraphs[0].add_run().add_picture(png, width=Cm(2))
    h.cell(0, 1).text = "TIMBRE SINTÉTICO — TESTES"
    f = sec.footer.add_table(rows=1, cols=2, width=Cm(17))
    f.cell(0, 1).text = "01/01/2099"
    for campo in ("PAGE", "NUMPAGES"):
        instr = OxmlElement("w:instrText")
        instr.text = campo
        instr.set(qn("xml:space"), "preserve")
        f.cell(0, 0).paragraphs[0].add_run()._r.append(instr)
    for label in ("Ordem de Serviço:", "Referência:", "Natureza:", "Investigado (s):",
                  "Vítima(s):", "Local:", "Data dos Fatos:", "Escrivão do feito:"):
        doc.add_paragraph(label)
    doc.add_paragraph("EXCELENTÍSSIMO (A) SENHOR (A)")
    for titulo, guia in (("RESUMO DOS FATOS", "{Resumir os fatos}"),
                         ("DILIGÊNCIAS REALIZADAS", "{Elencar diligências}"),
                         ("CONCLUSÃO", "{breve descrição da conclusão}")):
        p = doc.add_paragraph(titulo)
        p.runs[0].font.name, p.runs[0].font.size = "Arial", Pt(12)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        doc.add_paragraph(guia)
    doc.add_paragraph("[local, Estado]")
    png.seek(0)
    doc.add_paragraph().add_run().add_picture(png, width=Cm(4))
    # Texto público do bloco previsto pelos testes; a imagem acima é geométrica.
    doc.add_paragraph("Alan Douglas Silva")
    doc.add_paragraph("Investigador de Polícia")
    doc.add_paragraph("A(o) Excelentíssimo (a)")
    doc.add_paragraph("{Nome do Delegado}")
    doc.add_paragraph("Delegado (a) de Polícia Civil")
    doc.save(destino)
    return destino
