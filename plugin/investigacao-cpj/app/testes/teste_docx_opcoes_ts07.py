"""Opções reais do gerador em modelo/imagens/CSV inteiramente fictícios."""
import hashlib
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[4]
GERADOR = ROOT / 'plugin/investigacao-cpj/skills/relatorio-ip-fraude/scripts/gerar_docx.py'


class TesteOpcoesDocx(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='cpj-ts07-')
        self.addCleanup(self.temp.cleanup)
        self.ws = Path(self.temp.name)
        self.rel = self.ws / 'casos/OS-FAKE-2026/03-relatorios'
        self.rel.mkdir(parents=True)
        self.csv = self.rel.parent / '02-analise/fluxo-financeiro.csv'
        self.csv.parent.mkdir()
        self.csv.write_text('seq;data;valor;origem_titular;destino_titular;origem_banco;destino_banco;fonte_pag;fls\n1;10/10/2026;100,00;ORIGEM FICTICIA;DESTINO FICTICIO;BANCO A;BANCO B;2;2\n', encoding='utf-8-sig')
        self.assinatura = self.ws / 'assinatura-ficticia.png'
        self.conteudo = self.rel / 'conteudo-ficticio.png'
        self.fluxo = self.ws / 'fluxograma-ficticio.png'
        for p,cor in ((self.assinatura,'red'),(self.conteudo,'blue'),(self.fluxo,'green')):
            Image.new('RGB',(120,40),cor).save(p)
        self.modelo = self.ws / 'modelos/MODELO RELATORIO DE INVESTIGACAO - CPJ 2026.docx'
        self.modelo.parent.mkdir()
        d = Document()
        for s in ('RESUMO DOS FATOS','DILIGÊNCIAS REALIZADAS','CONCLUSÃO'):
            d.add_paragraph(s)
            d.add_paragraph({'RESUMO DOS FATOS':'{Resumir os fatos fictícios}','DILIGÊNCIAS REALIZADAS':'{Elencar as diligências fictícias}','CONCLUSÃO':'{breve descrição fictícia}'}[s])
        d.add_picture(str(self.assinatura))
        d.add_paragraph('SUBSCRITOR INTEIRAMENTE FICTÍCIO')
        d.save(self.modelo)
        self.minuta = self.rel / 'minuta-v01.md'
        self.minuta.write_text('---\nordem_servico: FAKE/2026\nreferencia: IPe fictício / Processo fictício\ndelegado_genero: F\n---\n## RESUMO DOS FATOS\nRelato fictício.\n## DILIGÊNCIAS REALIZADAS\nDocumento fictício examinado.\n## CONCLUSÃO\nSomente dados fictícios.\n',encoding='utf-8')

    def gerar(self,*opcoes,conteudo=False):
        if conteudo:
            texto=self.minuta.read_text(encoding='utf-8').replace('Documento fictício examinado.','Documento fictício examinado.\n\n![Imagem de conteúdo](conteudo-ficticio.png)')
            self.minuta.write_text(texto,encoding='utf-8')
        insumos=(self.modelo,self.csv,self.minuta,self.assinatura,self.conteudo,self.fluxo)
        hashes={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in insumos}
        saida=self.rel/'saida.docx'
        env={**os.environ,'CPJ_WORKSPACE':str(self.ws),'PYTHONUTF8':'1'}
        env.pop('CPJ_WORKSPACES',None)
        p=subprocess.run([sys.executable,str(GERADOR),str(self.minuta),'--saida',str(saida),'--modelo',str(self.modelo),*opcoes],cwd=self.ws,env=env,capture_output=True,text=True,encoding='utf-8',timeout=60)
        self.assertEqual(p.returncode,0,p.stdout+p.stderr)
        self.assertNotIn('Aviso ao gerar fluxograma',p.stderr)
        self.assertEqual(hashes,{q:hashlib.sha256(q.read_bytes()).hexdigest() for q in insumos},'insumos alterados')
        d = Document(saida)
        return [d.part.related_parts[x._inline.xpath('.//a:blip')[0].get(qn('r:embed'))].blob for x in d.inline_shapes]

    def test_csv_explicito_gera_fluxograma_e_preserva_assinatura_padrao(self):
        medias=self.gerar('--fluxo-csv',str(self.csv))
        self.assertEqual(len(medias),2)
        self.assertIn(self.assinatura.read_bytes(),medias)

    def test_csv_explicito_sem_assinatura_preserva_fluxograma(self):
        medias=self.gerar('--fluxo-csv',str(self.csv),'--sem-assinatura')
        self.assertEqual(len(medias),1)
        self.assertNotIn(self.assinatura.read_bytes(),medias)

    def test_descoberta_automatica_preserva_fluxograma_sem_assinatura(self):
        medias=self.gerar('--sem-assinatura')
        self.assertEqual(len(medias),1)

    def test_png_explicito_e_imagem_markdown_preservados_sem_assinatura(self):
        medias=self.gerar('--fluxograma',str(self.fluxo),'--sem-assinatura',conteudo=True)
        self.assertCountEqual(medias,[self.conteudo.read_bytes(),self.fluxo.read_bytes()])

    def test_assinatura_padrao_e_imagens_inseridas_preservadas(self):
        medias=self.gerar('--fluxograma',str(self.fluxo),conteudo=True)
        self.assertCountEqual(medias,[self.assinatura.read_bytes(),self.conteudo.read_bytes(),self.fluxo.read_bytes()])

    def test_sem_fluxograma_remove_so_assinatura_modelo(self):
        medias=self.gerar('--sem-fluxograma','--sem-assinatura',conteudo=True)
        self.assertEqual(medias,[self.conteudo.read_bytes()])

    def test_sem_fluxograma_mantem_assinatura_padrao(self):
        self.assertEqual(self.gerar('--sem-fluxograma'),[self.assinatura.read_bytes()])

    def test_imagem_inserida_com_mesmos_bytes_do_modelo_nao_e_removida(self):
        # A origem do parágrafo distingue assinatura de conteúdo, não o rId
        # nem a aparência da imagem (python-docx deduplica blobs iguais).
        self.conteudo.write_bytes(self.assinatura.read_bytes())
        self.assertEqual(self.gerar('--sem-fluxograma','--sem-assinatura',conteudo=True),[self.conteudo.read_bytes()])


if __name__=='__main__':
    unittest.main(verbosity=2)
