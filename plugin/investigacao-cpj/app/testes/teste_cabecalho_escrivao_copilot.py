#!/usr/bin/env python3
"""Garante que o escrivao da O.S. chegue ao cabecalho do DOCX CPJ."""
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[4]
MODELO = ROOT / "modelos" / "MODELO RELATORIO DE INVESTIGACAO - CPJ 2026.docx"
GERADOR = ROOT / "plugin" / "investigacao-cpj" / "skills" / "relatorio-ip-fraude" / "scripts" / "gerar_docx.py"


class CabecalhoEscrivao(unittest.TestCase):
    def gerar(self, meta_extra):
        pasta = Path(self.temp.name)
        minuta = pasta / f"minuta-{len(list(pasta.glob('minuta-*.md')))}.md"
        saida = minuta.with_suffix(".docx")
        meta = {
            "ordem_servico": "901/2026",
            "natureza": "Estelionato fictício",
            "data_fatos": "10/09/2026",
            "local_data": "Presidente Prudente, SP, 30 de setembro de 2026",
            "data_rodape": "30/09/2026",
            "delegado": "DELEGADO FICTÍCIO",
            "delegado_genero": "M",
        }
        meta.update(meta_extra)
        front = "\n".join(f"{k}: {v}" for k, v in meta.items())
        minuta.write_text(
            f"---\n{front}\n---\n\n"
            "## RESUMO DOS FATOS\n\nFato fictício para teste.\n\n"
            "## DILIGÊNCIAS REALIZADAS\n\nDiligência fictícia para teste.\n\n"
            "## CONCLUSÃO\n\nConclusão fictícia para teste.\n",
            encoding="utf-8",
        )
        resultado = subprocess.run(
            [sys.executable, str(GERADOR), str(minuta), "--saida", str(saida), "--modelo", str(MODELO)],
            capture_output=True,
            text=True,
            check=True,
            env={**os.environ, "PYTHONUTF8": "1"},
        )
        return Document(saida), resultado.stdout

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="cpj-escrivao-copilot-")

    def tearDown(self):
        self.temp.cleanup()

    def test_escrivao_fica_no_cabecalho_apos_data_dos_fatos(self):
        documento, _ = self.gerar({"escrivao": "ESCRIVÃO FICTÍCIO"})
        textos = [p.text for p in documento.paragraphs]
        linha = next(i for i, texto in enumerate(textos) if texto.startswith("Escrivão do feito:"))
        data = next(i for i, texto in enumerate(textos) if texto.startswith("Data dos Fatos:"))
        saudacao = next(i for i, texto in enumerate(textos) if texto.startswith("EXCELENTÍSSIMO"))
        self.assertIn("ESCRIVÃO FICTÍCIO", textos[linha])
        self.assertEqual(linha, data + 1)
        self.assertLess(linha, saudacao)

    def test_escrivao_ausente_fica_com_alerta_vermelho_sem_inferencia(self):
        documento, _ = self.gerar({})
        linha = next(p for p in documento.paragraphs if p.text.startswith("Escrivão do feito:"))
        self.assertIn("[OBTER: NOME DO ESCRIVÃO DO FEITO]", linha.text)
        run_alerta = next(r for r in linha.runs if "OBTER:" in r.text)
        self.assertEqual(run_alerta.font.color.rgb, RGBColor(255, 0, 0))


class ApiMinutaEscrivao(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="cpj-escrivao-api-")
        cls.ws = Path(cls.temp.name)
        cls.env_anterior = os.environ.get("CPJ_WORKSPACE")
        os.environ["CPJ_WORKSPACE"] = str(cls.ws)
        modelo = cls.ws / "casos" / "_MODELO-CASO"
        for pasta in ("00-originais", "01-extracao", "02-analise", "03-relatorios"):
            (modelo / pasta).mkdir(parents=True, exist_ok=True)
        (modelo / "caso.json").write_text(
            '{"datas": {}, "financeiro": {}, "ip": {}, "resultado": {}, "relatorios": []}',
            encoding="utf-8",
        )
        config = cls.ws / "config"
        config.mkdir(parents=True, exist_ok=True)
        (config / "solo.json").write_text('{"ativo": true}', encoding="utf-8")
        app_dir = ROOT / "plugin" / "investigacao-cpj" / "app"
        base_dir = ROOT / "plugin" / "investigacao-cpj" / "skills" / "base-cpj" / "scripts"
        sys.path.insert(0, str(app_dir))
        sys.path.insert(0, str(base_dir))
        import caso as C
        import servidor
        cls.C = C
        cls.servidor = servidor
        C.WS = str(cls.ws)
        C.CASOS = str(cls.ws / "casos")
        C.MODELO = str(modelo)
        servidor.WS = str(cls.ws)
        servidor.app.config.update(TESTING=True)
        C.novo(
            os_num="901/2026",
            id_="OS-901-2026",
            extras={"escrivao": "ESCRIVÃO FICTÍCIO"},
        )
        servidor.auth.salvar_usuario(
            "admin_teste", "Administrador Fictício", "admin", "SenhaFicticia123"
        )

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()
        if cls.env_anterior is None:
            os.environ.pop("CPJ_WORKSPACE", None)
        else:
            os.environ["CPJ_WORKSPACE"] = cls.env_anterior

    def test_minuta_nova_recebe_escrivao_salvo_na_os(self):
        resposta = self.servidor.app.test_client().get(
            "/api/casos/OS-901-2026/minuta", base_url="http://127.0.0.1"
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(resposta.get_json()["meta"]["escrivao"], "ESCRIVÃO FICTÍCIO")


if __name__ == "__main__":
    unittest.main(verbosity=2)
