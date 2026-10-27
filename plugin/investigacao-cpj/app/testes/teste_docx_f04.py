#!/usr/bin/env python3
"""Verifica preservação do modelo DOCX recebido; runner usa fixture sintética no clone."""
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fixtures_docx import criar_modelo

AQUI = Path(__file__).resolve().parent
ROOT = AQUI.parents[3]
GERADOR = ROOT / "plugin" / "investigacao-cpj" / "skills" / "relatorio-ip-fraude" / "scripts" / "gerar_docx.py"

MINUTA = """---
ordem_servico: 456/2026
referencia: IPe nº FICTÍCIO-456/2026 / Processo nº FICTÍCIO-789/2026
natureza: Estelionato (art. 171, CP)
investigados: INVESTIGADO FICTÍCIO
vitimas: VÍTIMA FICTÍCIA
local: Rua Fictícia, 10, Presidente Prudente/SP
data_fatos: 10/03/2026
local_data: Presidente Prudente, SP, 28 de setembro de 2026
data_rodape: 28/09/2026
delegado: DELEGADO FICTÍCIO
delegado_genero: M
---
## RESUMO DOS FATOS
Em 10/03/2026, a vítima fictícia relatou transferência de R$ 1.234,56, conforme (pág. 2 do PDF; fls. 14). O relato é fictício e serve apenas para validar a conversão do documento.
## DILIGÊNCIAS REALIZADAS
### Análise financeira
A transferência foi registrada em documento fictício, sem atribuição automática de autoria.
| Data | Histórico | Valor (R$) |
| --- | --- | ---: |
| 10/03/2026 | PIX FICTÍCIO | R$ 1.234,56 |
| 11/03/2026 | DEVOLUÇÃO FICTÍCIA | -R$ 24,10 |
- Conferência da página citada (pág. 2 do PDF; fls. 15).
## CONCLUSÃO
Os elementos fictícios documentam a transferência, sem permitir conclusão quanto à autoria.
"""


class TesteDocxF04(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="cpj-docx-f04-")
        pasta = Path(cls.temp.name)
        cls.modelo_path = criar_modelo(pasta / "modelos")
        modelo = Document(cls.modelo_path)
        for p in modelo.paragraphs:
            if p.text in {"RESUMO DOS FATOS", "DILIGÊNCIAS REALIZADAS", "CONCLUSÃO"}:
                p._p.addnext(OxmlElement("w:p"))
        modelo.save(cls.modelo_path)
        cls.saida = pasta / "relatorio-ficticio.docx"
        minuta = pasta / "minuta-ficticia.md"
        minuta.write_text(MINUTA, encoding="utf-8")
        subprocess.run(
            [sys.executable, str(GERADOR), str(minuta), "--saida", str(cls.saida), "--modelo", str(cls.modelo_path), "--sem-fluxograma"],
            check=True,
            capture_output=True,
            text=True,
            env={**os.environ, "CPJ_WORKSPACE": str(pasta), "CPJ_SEM_AGENTE_EMBUTIDO": "1", "PYTHONUTF8": "1"},
        )
        cls.modelo = Document(cls.modelo_path)
        cls.documento = Document(cls.saida)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_margens_timbre_rodape_e_paginacao_preservados(self):
        modelo_secao = self.modelo.sections[0]
        saida_secao = self.documento.sections[0]
        self.assertEqual(
            (saida_secao.page_width, saida_secao.page_height, saida_secao.top_margin,
             saida_secao.bottom_margin, saida_secao.left_margin, saida_secao.right_margin,
             saida_secao.header_distance, saida_secao.footer_distance),
            (modelo_secao.page_width, modelo_secao.page_height, modelo_secao.top_margin,
             modelo_secao.bottom_margin, modelo_secao.left_margin, modelo_secao.right_margin,
             modelo_secao.header_distance, modelo_secao.footer_distance),
        )
        self.assertEqual(saida_secao.header.tables[0].cell(0, 1).text, modelo_secao.header.tables[0].cell(0, 1).text)
        self.assertEqual(
            len(saida_secao.header._element.xpath(".//pic:pic")),
            len(modelo_secao.header._element.xpath(".//pic:pic")),
        )
        self.assertIn("28/09/2026", saida_secao.footer.tables[0].cell(0, 1).text)
        for documento in (self.modelo, self.documento):
            instrucoes = [x.text for x in documento.sections[0].footer._element.xpath(".//w:instrText")]
            self.assertIn("PAGE", instrucoes)
            self.assertIn("NUMPAGES", instrucoes)

    def test_tabela_tem_largura_definida_e_cabecalho_repetido(self):
        tabela = self.documento.tables[0]
        secao = self.documento.sections[0]
        largura_conteudo = secao.page_width - secao.left_margin - secao.right_margin
        largura_twips = int(round(largura_conteudo / 635))
        tbl_w = tabela._tbl.tblPr.find(qn("w:tblW"))
        largura_colunas = sum(int(col.get(qn("w:w"))) for col in tabela._tbl.tblGrid.gridCol_lst)
        self.assertEqual(tbl_w.get(qn("w:type")), "dxa")
        self.assertEqual(int(tbl_w.get(qn("w:w"))), largura_twips)
        self.assertEqual(largura_colunas, largura_twips)
        self.assertFalse(tabela.autofit)
        marcacao = tabela.rows[0]._tr.trPr.find(qn("w:tblHeader"))
        self.assertIsNotNone(marcacao)

    def test_valores_monetarios_sao_alinhados_e_tipografados(self):
        tabela = self.documento.tables[0]
        for linha in tabela.rows[1:]:
            celula = linha.cells[2]
            self.assertEqual(celula.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
            for run in celula.paragraphs[0].runs:
                if not run.text.strip():
                    continue
                self.assertEqual(run.font.name, "Arial")
                self.assertEqual(run.font.size.pt, 9)

    def test_assinatura_nome_cargo_formam_bloco_sem_quebra_forcada(self):
        paragrafos_assinatura = [
            p for p in self.documento.paragraphs if p._p.xpath(".//pic:pic")
        ]
        self.assertEqual(len(paragrafos_assinatura), 1)
        assinatura = paragrafos_assinatura[0]
        self.assertFalse(assinatura.paragraph_format.page_break_before)
        self.assertTrue(assinatura.paragraph_format.keep_with_next)
        self.assertTrue(assinatura.paragraph_format.keep_together)
        textos = [p.text for p in self.documento.paragraphs]
        self.assertIn("ALAN DOUGLAS SILVA", textos)
        self.assertIn("Investigador de Polícia", textos)
        self.assertIn("DELEGADO FICTÍCIO", textos)
        indice = next(i for i, p in enumerate(self.documento.paragraphs) if p._p == assinatura._p)
        bloco = self.documento.paragraphs[indice:indice + 3]
        self.assertEqual([p.text for p in bloco[1:]], ["ALAN DOUGLAS SILVA", "Investigador de Polícia"])
        self.assertTrue(bloco[1].paragraph_format.keep_with_next)
        for p in bloco:
            self.assertTrue(p.paragraph_format.keep_together)
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(p.paragraph_format.first_line_indent, 0)
            self.assertFalse(p.paragraph_format.page_break_before)
        self.assertEqual(
            len(self.documento._element.body.xpath(".//pic:pic")),
            len(self.modelo._element.body.xpath(".//pic:pic")),
        )

    def test_citacoes_lista_e_estilos_do_modelo_sao_preservados(self):
        textos = [p.text for p in self.documento.paragraphs]
        self.assertTrue(any("(Fls. 14)" in texto for texto in textos))
        self.assertTrue(any("(Fls. 15)" in texto for texto in textos))
        self.assertNotIn("do PDF", "\n".join(textos))
        self.assertTrue(any(texto.startswith("- Conferência") for texto in textos))
        for nome in ("RESUMO DOS FATOS", "DILIGÊNCIAS REALIZADAS", "CONCLUSÃO"):
            modelo = next(p for p in self.modelo.paragraphs if p.text == nome)
            saida = next(p for p in self.documento.paragraphs if p.text == nome)
            self.assertEqual(saida.style.name, modelo.style.name)
            self.assertEqual(saida.alignment, modelo.alignment)
            self.assertEqual(saida.runs[0].font.name, modelo.runs[0].font.name)
            self.assertEqual(saida.runs[0].font.size, modelo.runs[0].font.size)
            self.assertEqual(saida.paragraph_format.space_before, modelo.paragraph_format.space_before)
            self.assertEqual(saida.paragraph_format.space_after, modelo.paragraph_format.space_after)


    def test_titulos_de_secao_permanecem_com_o_primeiro_paragrafo(self):
        nomes = {"RESUMO DOS FATOS", "DILIGÊNCIAS REALIZADAS", "CONCLUSÃO"}
        titulos = [p for p in self.documento.paragraphs if p.text in nomes]
        self.assertEqual(len(titulos), 3)
        for titulo in titulos:
            with self.subTest(secao=titulo.text):
                self.assertTrue(titulo.paragraph_format.keep_with_next)
                self.assertTrue(titulo.paragraph_format.keep_together)


    def test_espacador_do_modelo_nao_separa_titulo_do_texto(self):
        nomes = {"RESUMO DOS FATOS", "DILIGÊNCIAS REALIZADAS", "CONCLUSÃO"}
        pars = self.documento.paragraphs
        for i, p in enumerate(pars):
            if p.text in nomes:
                with self.subTest(secao=p.text):
                    self.assertEqual(pars[i + 1].text, "")
                    self.assertTrue(pars[i + 1].paragraph_format.keep_with_next)

if __name__ == "__main__":
    unittest.main(verbosity=2)
