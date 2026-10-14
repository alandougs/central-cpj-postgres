"""SIG01: citações e assinatura no DOCX; somente modelo geométrico fictício."""
import hashlib
import os
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from docx import Document
import teste_docx_opcoes_ts07 as opcoes


class FolhasAssinatura(opcoes.TesteOpcoesDocx):
    def test_citacao_dupla_mostra_so_folhas_sem_mudar_markdown(self):
        s=self.minuta.read_text(encoding='utf-8').replace('Documento fictício examinado.','Documento fictício (págs. 242–243 do PDF; fls. 12–13). Outro (pág. 9 do PDF; fls. 27).\n\n| Origem | Fonte |\n|---|---|\n| Fictícia | (págs. 20–21 do PDF; fls. 30–31) |')
        self.minuta.write_text(s,encoding='utf-8')
        sha=hashlib.sha256(self.minuta.read_bytes()).hexdigest()
        self.gerar('--sem-fluxograma')
        d=Document(self.rel/'saida.docx')
        texto='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
        self.assertIn('(fls. 12–13)',texto)
        self.assertIn('(fls. 27)',texto)
        self.assertIn('(fls. 30–31)',texto)
        self.assertNotIn('do PDF',texto)
        self.assertEqual(sha,hashlib.sha256(self.minuta.read_bytes()).hexdigest())

    def test_so_pagina_pdf_nao_vira_folha_inventada(self):
        self.minuta.write_text(self.minuta.read_text(encoding='utf-8').replace('Documento fictício examinado.','Documento (pág. 19 do PDF), sem folha informada.'),encoding='utf-8')
        self.gerar('--sem-fluxograma')
        texto='\n'.join(p.text for p in Document(self.rel/'saida.docx').paragraphs)
        self.assertIn('(pág. 19 do PDF)',texto)
        self.assertNotIn('fls. 19',texto)

    def test_assinatura_incluida_junto_ao_subscritor_sem_pagina_forcada(self):
        medias=self.gerar('--sem-fluxograma')
        self.assertIn(self.assinatura.read_bytes(),medias)
        d=Document(self.rel/'saida.docx')
        p=next(p for p in d.paragraphs if p._element.xpath('.//pic:pic'))
        self.assertFalse(p.paragraph_format.page_break_before)
        self.assertTrue(p.paragraph_format.keep_with_next)

    def test_final_da_etapa_regenera_docx_do_modelo_mesmo_se_agente_entregou_arquivo_recente(self):
        destino=self.rel/'RELATORIO-OS-FAKE-2026-v01.docx'
        d=Document();d.add_paragraph('Arquivo fictício do agente, sem assinatura nem modelo');d.save(destino)
        sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
        import plantao
        with patch.dict(os.environ,{'CPJ_WORKSPACE':str(self.ws)}):
            resultado=plantao.pos_processar(str(self.ws),'OS-FAKE-2026','relatorio')
        self.assertEqual(resultado['docx'],destino.name)
        self.assertGreater(len(Document(destino).inline_shapes),0,'Pós-processamento aceitou DOCX sem assinatura')


if __name__=='__main__':unittest.main(verbosity=2)
