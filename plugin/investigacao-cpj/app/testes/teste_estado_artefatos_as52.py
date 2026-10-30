"""AS52: regressões de estado e artefatos, somente fixtures temporárias."""
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import zipfile

APP = Path(__file__).resolve().parents[1]


class EstadoArtefatosAS52(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="cpj-as52-")
        cls.ws = Path(cls.temp.name)
        cls.env = mock.patch.dict(os.environ, CPJ_WORKSPACE=str(cls.ws), CPJ_SEM_AGENTE_EMBUTIDO="1")
        cls.env.start()
        sys.path.insert(0, str(APP))
        from rotas import comum, casos, relatorios
        import tarefas
        from flask import Flask, request
        cls.comum, cls.casos, cls.rel, cls.tarefas = comum, casos, relatorios, tarefas
        cls.C = comum.C
        assert Path(cls.C.WS).resolve() == cls.ws.resolve()
        cls.app = Flask(__name__)
        cls.app.add_url_rule("/caso/<id_>", view_func=casos.api_caso_salvar.__wrapped__, methods=["POST"])
        cls.app.add_url_rule("/final/<id_>", view_func=relatorios.api_final.__wrapped__, methods=["POST"])
        @cls.app.before_request
        def usuario():
            request.usuario = {"login": "ficticio", "perfil": "admin"}

    @classmethod
    def tearDownClass(cls):
        cls.env.stop()
        cls.temp.cleanup()

    def setUp(self):
        self.id = "OS-" + self._testMethodName.replace("_", "-") + "-2099"
        self.pasta = self.ws / "casos" / self.id
        self.relatorio = self.pasta / "03-relatorios"
        self.relatorio.mkdir(parents=True)
        self.C.salvar({"id": self.id, "status": "recebido", "datas": {}, "visto": True,
                       "preenchimento_os": {"fontes": {"natureza": {"documento": "ficticio.md"}},
                                             "conflitos": {}, "pendentes": []}})
        self.client = self.app.test_client()
        self.audit = mock.patch.multiple(self.casos, auditar=lambda *a: None)
        self.audit.start()
        self.addCleanup(self.audit.stop)

    def post(self, dados):
        return self.client.post("/caso/" + self.id, json=dados)

    def test_ficha_sem_revisao_nao_desfaz_entrega(self):
        self.C.status(self.id, "entregue", arquivo="ficticio.docx")
        antes = self.C.carregar(self.id)
        r = self.post({"origem_edicao": "ficha", "natureza": "antiga", "status": "recebido"})
        self.assertEqual(r.status_code, 409)
        self.assertEqual(self.C.carregar(self.id), antes)

    def test_ficha_cas_obsoleta_preserva_todo_estado(self):
        antiga = self.C.carregar(self.id)["revisao"]
        self.C.status(self.id, "entregue", arquivo="ficticio.docx")
        antes = self.C.carregar(self.id)
        self.assertEqual(self.post({"origem_edicao": "ficha", "revisao": antiga,
                                   "natureza": "nova", "status": "recebido"}).status_code, 409)
        self.assertEqual(self.C.carregar(self.id), antes)

    def test_campos_proveniencia_status_uma_unica_revisao(self):
        antes = self.C.carregar(self.id)["revisao"]
        r = self.post({"origem_edicao": "ficha", "revisao": antes, "natureza": "nova", "status": "minuta"})
        self.assertEqual(r.status_code, 200)
        c = self.C.carregar(self.id)
        self.assertEqual(c["revisao"], antes + 1)
        self.assertEqual(c["status"], "minuta")
        self.assertNotIn("natureza", c["preenchimento_os"]["fontes"])

    def test_resposta_cas_corresponde_a_propria_gravacao(self):
        antes = self.C.carregar(self.id)["revisao"]
        with mock.patch.object(self.casos, "auditar", side_effect=lambda *a: self.C.set_campos(self.id, {"natureza": "edicao concorrente"})):
            r = self.post({"origem_edicao": "ficha", "revisao": antes, "natureza": "minha edicao"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json["revisao"], antes + 1)
        self.assertEqual(r.json["natureza"], "minha edicao")
        self.assertEqual(self.C.carregar(self.id)["revisao"], antes + 2)

    def test_status_invalido_nao_publica_campos_parciais(self):
        antes = self.C.carregar(self.id)
        self.assertEqual(self.post({"revisao": antes["revisao"], "natureza": "nova", "status": "invalido"}).status_code, 400)
        self.assertEqual(self.C.carregar(self.id), antes)

    def test_revisao_invalida_nao_vira_escrita_sem_cas(self):
        antes = self.C.carregar(self.id)
        for revisao in ("errada", True, [], 1.5, None):
            with self.subTest(revisao=revisao):
                self.assertIn(self.post({"origem_edicao": "ficha", "revisao": revisao, "natureza": "nova"}).status_code, (400, 409))
                self.assertEqual(self.C.carregar(self.id), antes)

    def test_patch_legado_preserva_baixa_e_status_isolado_continua(self):
        self.C.status(self.id, "entregue", arquivo="ficticio.docx")
        baixa = self.C.carregar(self.id)["baixa"]
        self.assertEqual(self.post({"observacoes": "anotação"}).status_code, 200)
        self.assertEqual(self.C.carregar(self.id)["baixa"], baixa)
        self.assertEqual(self.post({"status": "minuta"}).status_code, 200)
        self.assertIsNone(self.C.carregar(self.id)["baixa"])

    def test_legado_misto_nao_desfaz_baixa(self):
        self.C.status(self.id, "entregue", arquivo="ficticio.docx")
        self.assertEqual(self.post({"natureza": "antiga", "status": "recebido"}).status_code, 409)
        self.assertEqual(self.C.carregar(self.id)["status"], "entregue")

    def pdf(self, codigo, criar=False):
        docx = self.relatorio / "ficticio.docx"
        docx.write_bytes(b"docx sintetico")
        pdf = docx.with_suffix(".pdf")
        pdf.write_bytes(b"%PDF-antigo")
        def converter(cmd, **kw):
            if criar:
                (Path(cmd[cmd.index("--outdir") + 1]) / pdf.name).write_bytes(b"%PDF-novo")
            return subprocess.CompletedProcess(cmd, codigo, "", "conversao falhou")
        with mock.patch.object(self.tarefas, "soffice", return_value="soffice-ficticio"), mock.patch.object(self.tarefas.subprocess, "run", side_effect=converter):
            return pdf, lambda: self.tarefas.gerar_pdf(str(self.ws), self.id, docx.name)

    def test_pdf_antigo_nao_aprova_conversao_falha(self):
        docx = self.relatorio / "ficticio.docx"
        docx.write_bytes(b"sintetico")
        pdf = docx.with_suffix(".pdf")
        pdf.write_bytes(b"%PDF-antigo")
        for codigo in (1, 0):
            with self.subTest(codigo=codigo), mock.patch.object(self.tarefas, "soffice", return_value="soffice-ficticio"), mock.patch.object(self.tarefas.subprocess, "run", return_value=subprocess.CompletedProcess([], codigo, "", "falhou")):
                with self.assertRaises(RuntimeError): self.tarefas.gerar_pdf(str(self.ws), self.id, docx.name)
                self.assertEqual(pdf.read_bytes(), b"%PDF-antigo")

    def test_pdf_publica_novo_so_apos_conversao(self):
        docx = self.relatorio / "ficticio.docx"
        docx.write_bytes(b"sintetico")
        pdf = docx.with_suffix(".pdf")
        pdf.write_bytes(b"%PDF-antigo")
        def converter(cmd, **kw):
            self.assertEqual(pdf.read_bytes(), b"%PDF-antigo")
            (Path(cmd[cmd.index("--outdir") + 1]) / pdf.name).write_bytes(b"%PDF-novo")
            return subprocess.CompletedProcess(cmd, 0, "", "")
        with mock.patch.object(self.tarefas, "soffice", return_value="soffice-ficticio"), mock.patch.object(self.tarefas.subprocess, "run", side_effect=converter):
            self.assertEqual(self.tarefas.gerar_pdf(str(self.ws), self.id, docx.name), {"pdf": pdf.name})
        self.assertEqual(pdf.read_bytes(), b"%PDF-novo")

    def importar(self, arquivos):
        fluxo = io.BytesIO()
        manifestados = []
        with zipfile.ZipFile(fluxo, "w") as z:
            for nome, valor in arquivos.items():
                b = json.dumps(valor).encode()
                caminho = nome if nome.startswith("casos/") else "casos/OS-IMPORTADO-2099/" + nome
                z.writestr("dados/" + caminho, b)
                manifestados.append({"caminho": caminho, "tamanho": len(b), "sha256": hashlib.sha256(b).hexdigest()})
            z.writestr("manifest.json", json.dumps({"schema": self.tarefas.EXPORT_SCHEMA, "arquivos": manifestados}))
        fluxo.seek(0)
        t = self.tarefas.Tarefas.__new__(self.tarefas.Tarefas)
        t.ws, t.at, t.indexar = str(self.pasta), lambda *a, **k: None, lambda: None
        return t._importar_zip("sintetico", fluxo, str(self.pasta / "staging"), lambda: None)

    def test_importacao_rejeita_tipos_inseguros_antes_de_publicar(self):
        for nome, valor in (("caso.json", []), ("caso.json", {"documentos": [{"paginas": "<img src=x onerror=alert(1)>"}]}),
                            ("caso.json", {"vitimas": "texto"}), ("caso.json", {"preenchimento_os": {"fontes": {"bo": {"pagina": "<svg onload=alert(1)>"}}}}),
                            ("processamento.json", {"trabalhos": ["erro"]}), ("processamento.json", {"trabalhos": [{"progresso": "<img>"}]})):
            with self.subTest(nome=nome, valor=valor):
                arquivos = {"caso.json": {}}
                arquivos[nome] = valor
                with self.assertRaises(ValueError): self.importar(arquivos)
                self.assertFalse((self.pasta / "casos/OS-IMPORTADO-2099").exists())

    def test_importacao_preserva_pacote_legado_com_campos_ausentes(self):
        resultado = self.importar({"caso.json": {"ficticio": True}, "processamento.json": {"trabalhos": []}})
        self.assertEqual(resultado["casos_novos"], ["OS-IMPORTADO-2099"])

    def test_importacao_id_divergente_nao_publica_nem_caso_valido(self):
        with self.assertRaises(ValueError):
            self.importar({"casos/OS-VALIDO-2099/caso.json": {"id": "OS-VALIDO-2099"},
                           "casos/OS-INVALIDO-2099/caso.json": {"id": "OS-OUTRO-2099"}})
        self.assertFalse((self.pasta / "casos/OS-VALIDO-2099").exists())

    def preparar_docx(self, vincular=True):
        from docx import Document
        minuta = self.relatorio / "minuta-v01.md"
        minuta.write_text("## CONCLUSÃO\nTexto sintético.", encoding="utf-8")
        docx = self.relatorio / ("RELATORIO-" + self.id + "-v01.docx")
        doc = Document()
        doc.add_paragraph("Texto sintético.")
        doc.save(docx)
        if vincular:
            from vinculos_docx import emitir
            registro = emitir(str(self.ws), self.id, minuta.name, docx.name,
                              hashlib.sha256(minuta.read_bytes()).hexdigest(),
                              hashlib.sha256(docx.read_bytes()).hexdigest(), "2099-01-01T00:00:00")
            self.C.atualizar(self.id, lambda c: c.update(vinculos_docx=[registro]))
        return minuta, docx

    def finalizar(self, docx, dados=None, durante=None):
        def conferir(*a):
            sha = hashlib.sha256((self.relatorio / "minuta-v01.md").read_bytes()).hexdigest()
            if durante:
                durante()
            return {"aprovado": True, "achados": [], "sha256_minuta": sha}
        with mock.patch.object(self.rel, "conferir_minuta_gate", side_effect=conferir), mock.patch.object(self.rel, "auditar"), mock.patch.object(self.rel, "auth") as auth, mock.patch.object(self.rel.threading, "Thread"):
            auth.obter.return_value = None
            return self.client.post("/final/" + self.id, json={"docx": docx.name, **(dados or {})})

    def test_final_bloqueia_docx_editado_mesmo_com_force_generico(self):
        minuta, docx = self.preparar_docx()
        docx.write_bytes(docx.read_bytes() + b"edicao Word")
        for dados in ({}, {"forcar": True, "justificativa": "Conferi os dados e assumo ressalvas."}):
            r = self.finalizar(docx, dados)
            self.assertEqual(r.status_code, 409)
            self.assertEqual(self.C.carregar(self.id)["status"], "recebido")
            self.assertFalse(list(self.relatorio.glob("*-FINAL.*")))

    def test_final_bloqueia_docx_legado_sem_vinculo(self):
        _, docx = self.preparar_docx(vincular=False)
        self.assertEqual(self.finalizar(docx).status_code, 409)

    def test_pacote_com_vinculo_fabricado_nao_autoriza_final(self):
        minuta, docx = self.preparar_docx(vincular=False)
        importado = "OS-VINCULO-FABRICADO-2099"
        nome_docx = "RELATORIO-" + importado + "-v01.docx"
        registro = {"schema": "cpj-vinculo-docx/1", "caso": importado, "minuta": minuta.name,
                    "docx": nome_docx, "sha256_minuta": hashlib.sha256(minuta.read_bytes()).hexdigest(),
                    "sha256_docx": hashlib.sha256(docx.read_bytes()).hexdigest(), "gerado_em": "2099-01-01T00:00:00"}
        caso = {"id": importado, "status": "recebido", "datas": {}, "vinculos_docx": [registro]}
        caso_bytes = json.dumps(caso).encode("utf-8")
        arquivos = {"caso.json": caso_bytes, "03-relatorios/minuta-v01.md": minuta.read_bytes(),
                    "03-relatorios/" + nome_docx: docx.read_bytes()}
        fluxo, manifestados = io.BytesIO(), []
        with zipfile.ZipFile(fluxo, "w") as z:
            for nome, conteudo in arquivos.items():
                caminho = "casos/" + importado + "/" + nome
                z.writestr("dados/" + caminho, conteudo)
                manifestados.append({"caminho": caminho, "tamanho": len(conteudo), "sha256": hashlib.sha256(conteudo).hexdigest()})
            z.writestr("manifest.json", json.dumps({"schema": self.tarefas.EXPORT_SCHEMA, "arquivos": manifestados}))
        fluxo.seek(0)
        t = self.tarefas.Tarefas.__new__(self.tarefas.Tarefas)
        t.ws, t.at, t.indexar = str(self.ws), lambda *a, **k: None, lambda: None
        t._importar_zip("sintetico", fluxo, str(self.pasta / "staging"), lambda: None)
        self.id = importado
        self.relatorio = self.ws / "casos" / importado / "03-relatorios"
        self.assertEqual(self.finalizar(self.relatorio / nome_docx).status_code, 409)
        self.assertEqual((self.relatorio.parent / "caso.json").read_bytes(), caso_bytes)

    def test_final_rejeita_mac_adulterado(self):
        _, docx = self.preparar_docx()
        self.C.atualizar(self.id, lambda c: c["vinculos_docx"][0].update(hmac_sha256="0" * 64))
        self.assertEqual(self.finalizar(docx).status_code, 409)

    def test_final_rejeita_vinculo_transferido_para_outro_caso(self):
        _, docx = self.preparar_docx()
        self.C.atualizar(self.id, lambda c: c["vinculos_docx"][0].update(caso="OS-OUTRO-CASO-2099"))
        self.assertEqual(self.finalizar(docx).status_code, 409)

    def test_final_publica_snapshot_vinculado(self):
        _, docx = self.preparar_docx()
        conteudo = docx.read_bytes()
        r = self.finalizar(docx)
        self.assertEqual(r.status_code, 200)
        self.assertEqual((self.relatorio / r.json["final"]).read_bytes(), conteudo)
        self.assertEqual(self.C.carregar(self.id)["status"], "entregue")

    def test_final_exige_aceite_especifico_do_hash_editado(self):
        _, docx = self.preparar_docx()
        from docx import Document
        doc = Document(docx)
        doc.add_paragraph("Edição humana conferida.")
        doc.save(docx)
        sha = hashlib.sha256(docx.read_bytes()).hexdigest()
        dados = {"forcar": True, "justificativa": "Conferi o documento editado no Word.", "confirmar_docx_sha256": "0" * 64}
        self.assertEqual(self.finalizar(docx, dados).status_code, 409)
        dados["confirmar_docx_sha256"] = sha
        r = self.finalizar(docx, dados)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(hashlib.sha256((self.relatorio / r.json["final"]).read_bytes()).hexdigest(), sha)
        self.assertEqual(self.C.carregar(self.id)["baixa"].get("sha256_docx"), sha)

    def test_final_detecta_troca_docx_durante_conferencia(self):
        _, docx = self.preparar_docx()
        r = self.finalizar(docx, durante=lambda: docx.write_bytes(b"trocado durante conferencia"))
        self.assertEqual(r.status_code, 409)
        self.assertFalse(list(self.relatorio.glob("*-FINAL.*")))

    def test_final_detecta_troca_minuta_durante_conferencia(self):
        minuta, docx = self.preparar_docx()
        r = self.finalizar(docx, durante=lambda: minuta.write_text("minuta trocada", encoding="utf-8"))
        self.assertEqual(r.status_code, 409)
        self.assertFalse(list(self.relatorio.glob("*-FINAL.*")))

    def test_geracao_docx_vincula_hashes_da_mesma_execucao(self):
        minuta, docx = self.preparar_docx(vincular=False)
        def gerar(cmd, **kw):
            from docx import Document
            novo = Document()
            novo.add_paragraph("gerado sintético")
            novo.save(cmd[cmd.index("--saida") + 1])
            return subprocess.CompletedProcess(cmd, 0, "", "")
        with mock.patch.object(self.tarefas.subprocess, "run", side_effect=gerar):
            self.tarefas.gerar_docx(str(self.ws), self.id, minuta.name)
        vinculo = self.C.carregar(self.id).get("vinculos_docx", [])
        self.assertTrue(vinculo)
        self.assertEqual(vinculo[-1]["sha256_docx"], hashlib.sha256(docx.read_bytes()).hexdigest())
        self.assertEqual(vinculo[-1]["sha256_minuta"], hashlib.sha256(minuta.read_bytes()).hexdigest())

    def test_geracao_docx_timeout_preserva_anterior(self):
        minuta, docx = self.preparar_docx()
        antes = docx.read_bytes()
        with mock.patch.object(self.tarefas.subprocess, "run", side_effect=subprocess.TimeoutExpired("sintetico", 180)):
            with self.assertRaises(RuntimeError): self.tarefas.gerar_docx(str(self.ws), self.id, minuta.name)
        self.assertEqual(docx.read_bytes(), antes)
        self.assertFalse(list(self.relatorio.glob(".docx-geracao-*")))

    def test_geracao_docx_usa_snapshot_da_minuta_contra_aba(self):
        minuta, docx = self.preparar_docx()
        original = minuta.read_bytes()
        def gerar(cmd, **kw):
            from docx import Document
            minuta.write_text("Texto B transitório.", encoding="utf-8")
            fonte = Path(cmd[2])
            self.assertEqual(fonte.parent, self.relatorio)
            novo = Document()
            novo.add_paragraph(fonte.read_text(encoding="utf-8"))
            novo.save(cmd[cmd.index("--saida") + 1])
            minuta.write_bytes(original)
            return subprocess.CompletedProcess(cmd, 0, "", "")
        with mock.patch.object(self.tarefas.subprocess, "run", side_effect=gerar):
            self.tarefas.gerar_docx(str(self.ws), self.id, minuta.name)
        from docx import Document
        texto = "\n".join(p.text for p in Document(docx).paragraphs)
        self.assertIn(original.decode("utf-8").replace("\r\n", "\n"), texto)
        self.assertNotIn("Texto B transitório", texto)
        self.assertFalse(list(self.relatorio.glob(".minuta-fonte-*")))

    def test_geracao_falha_vinculo_preserva_historico_anterior(self):
        minuta, docx = self.preparar_docx()
        antes = docx.read_bytes()
        def gerar(cmd, **kw):
            from docx import Document
            novo = Document()
            novo.add_paragraph("nova geracao sintetica")
            novo.save(cmd[cmd.index("--saida") + 1])
            return subprocess.CompletedProcess(cmd, 0, "", "")
        with mock.patch.object(self.tarefas.subprocess, "run", side_effect=gerar), mock.patch.object(self.C, "atualizar", side_effect=ValueError("falha sintética")):
            with self.assertRaises(RuntimeError): self.tarefas.gerar_docx(str(self.ws), self.id, minuta.name)
        self.assertIn(antes, [p.read_bytes() for p in (self.relatorio / "historico-docx").iterdir()])
        self.assertEqual(docx.read_bytes(), antes)


if __name__ == "__main__":
    unittest.main(verbosity=2)
