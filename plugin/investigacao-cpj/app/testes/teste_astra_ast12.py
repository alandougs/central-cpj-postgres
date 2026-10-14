"""AST12: editor real no Edge e HTTP, somente workspace e documentos fictícios."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import time
import unittest

APP = Path(__file__).resolve().parents[1]
RUNTIME = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node"


class EditorAST12(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="cpj-ast12-")
        cls.ws = Path(cls.temp.name)
        cls.old_env = {k: os.environ.get(k) for k in ("CPJ_WORKSPACE", "CPJ_SEM_AGENTE_EMBUTIDO")}
        os.environ.update(CPJ_WORKSPACE=str(cls.ws), CPJ_SEM_AGENTE_EMBUTIDO="1")
        model = cls.ws / "casos/_MODELO-CASO"
        model.mkdir(parents=True)
        (model / "caso.json").write_text(json.dumps({
            "datas": {}, "financeiro": {}, "ip": {}, "resultado": {}, "relatorios": []
        }), encoding="utf-8")
        from docx import Document
        doc = Document()
        for text in ("Ordem de Serviço:", "Referência:", "EXCELENTÍSSIMO SENHOR DOUTOR DELEGADO DE POLÍCIA",
                     "RESUMO DOS FATOS", "{Resumir os fatos}", "DILIGÊNCIAS REALIZADAS",
                     "{Elencar as diligências}", "CONCLUSÃO", "{breve descrição}"):
            doc.add_paragraph(text)
        (cls.ws / "modelos").mkdir()
        doc.save(cls.ws / "modelos/MODELO RELATORIO DE INVESTIGACAO - CPJ 2026.docx")
        sys.path.insert(0, str(APP))
        import servidor
        cls.s = servidor
        cls.s.auth.salvar_usuario("admin_ast12", "Pessoa fictícia", "admin", "SenhaFicticia123")
        (cls.ws / "config/solo.json").write_text('{"ativo":true}', encoding="utf-8")
        cls.control = {"fail": False, "delay": 0, "delay_post": 0, "fail_get": False}
        from flask import request, jsonify

        @servidor.app.before_request
        def controlled_failure():
            if request.path == "/_teste_ast12" and request.method == "POST":
                cls.control.update(request.get_json())
                return jsonify(ok=True)
            if request.path.endswith("/minuta"):
                if request.method == "GET" and cls.control["fail_get"]:
                    cls.control["fail_get"] = False
                    return jsonify(erro="Falha de leitura fictícia"), 503
                if request.method == "POST" and cls.control["fail"]:
                    cls.control["fail"] = False
                    return jsonify(erro="Falha temporária fictícia"), 503
                if request.method == "POST" and cls.control["delay_post"]:
                    delay = cls.control["delay_post"]
                    cls.control["delay_post"] = 0
                    time.sleep(delay)
                if request.method == "GET" and cls.control["delay"]:
                    delay = cls.control["delay"]
                    cls.control["delay"] = 0
                    time.sleep(delay)

        from werkzeug.serving import make_server
        cls.server = make_server("127.0.0.1", 0, servidor.app, threaded=True)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = f"http://127.0.0.1:{cls.server.server_port}"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.thread.join(timeout=5)
        cls.temp.cleanup()
        for key, value in cls.old_env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def setUp(self):
        self.__class__.number = getattr(self.__class__, "number", 10) + 2
        self.cid = self.s.C.novo(f"{self.number}/2099")["id"]
        self.cid2 = self.s.C.novo(f"{self.number + 1}/2099")["id"]
        self.control.update(fail=False, delay=0, delay_post=0, fail_get=False)

    def browser(self, body):
        script = r"""
const {chromium}=require(process.env.CPJ_NODE_MODULES+'/playwright');
const assert=require('node:assert/strict');
(async()=>{
 const browser=await chromium.launch({channel:'msedge',headless:true});
 try{
 const context=await browser.newContext();const page=await context.newPage();
 const base=process.env.AST12_URL,cid=process.env.AST12_CID,cid2=process.env.AST12_CID2;
 await page.goto(base);await page.waitForFunction(()=>!document.querySelector('#app').hidden);
 const open=async(id=cid,arq)=>{await page.evaluate(([id,arq])=>abrirEditor(id,arq),[id,arq]);await page.locator('#ed-secoes textarea').first().waitFor()};
 const text=()=>page.locator('#ed-secoes textarea[data-s="RESUMO DOS FATOS"]');
 const post=async(id,data)=>{const r=await context.request.post(`${base}/api/casos/${id}/minuta`,{headers:{'X-CPJ':'1'},data});return {status:r.status(),data:await r.json()}};
 const state=async(s)=>page.waitForFunction(s=>document.querySelector('#editor').dataset.estado===s,s);
 """ + body + r"""
 }finally{await browser.close()}
})().catch(e=>{console.error(e);process.exit(1)});
"""
        js = self.ws / "browser-test.cjs"
        js.write_text(script, encoding="utf-8")
        env = dict(os.environ, CPJ_NODE_MODULES=str(RUNTIME / "node_modules"), AST12_URL=self.base,
                   AST12_CID=self.cid, AST12_CID2=self.cid2)
        result = subprocess.run([str(RUNTIME / "bin/node.exe"), str(js)], env=env,
                                capture_output=True, text=True, encoding="utf-8", timeout=95)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_salvar_nao_gera_docx_autosave_preserva_secoes_e_docx_explicito(self):
        # Acoplar DOCX ao save ou fechar o editor após salvar perde este contrato.
        self.browser(r"""
 await page.evaluate(async id=>{abrirAba('casos');await abrirCaso(id)},cid);
 await page.locator('#b-nova-minuta').click();await text().fill('Rascunho fictício A');
 await page.locator('#ed-meta input[data-m="delegado_genero"]').fill('M');
 await page.locator('#ed-secoes textarea[data-s="DILIGÊNCIAS REALIZADAS"]').fill('Diligência fictícia documentada.');
 await page.locator('#ed-secoes textarea[data-s="CONCLUSÃO"]').fill('Conclusão fictícia.');
 const save=page.waitForResponse(r=>r.url().endsWith('/minuta')&&r.request().method()==='POST');
 await page.locator('#ed-salvar').click();await save;
 assert.equal(await page.locator('#editor').isVisible(),true);
 let r=await context.request.get(`${base}/api/casos/${cid}`);let j=await r.json();
 assert.equal(j.arquivos.filter(x=>/\.docx$/i.test(x.caminho)).length,0);
 await text().fill('Rascunho fictício B');await state('salvo');
 r=await context.request.get(`${base}/api/casos/${cid}/minuta`);j=await r.json();
 assert.equal(j.versao_num,2);assert.equal(j.secoes['RESUMO DOS FATOS'],'Rascunho fictício B');
 let docx=page.waitForResponse(r=>r.url().endsWith('/docx'));await page.locator('#ed-docx').click();let dr=await docx;assert.equal(dr.status(),200,await dr.text());await page.waitForFunction(()=>document.querySelector('#ed-msg').textContent.includes('DOCX gerado'));
 await page.locator('#ed-cancelar').click();await open(cid,'minuta-v01.md');
 assert.equal(await text().inputValue(),'Rascunho fictício A');
 docx=page.waitForResponse(r=>r.url().endsWith('/docx'));await page.locator('#ed-docx').click();dr=await docx;assert.equal(dr.status(),200,await dr.text());await page.waitForFunction(()=>document.querySelector('#ed-msg').textContent.includes('DOCX gerado'));
 const final=await context.request.post(`${base}/api/casos/${cid}/final`,{headers:{'X-CPJ':'1'},data:{docx:(await dr.json()).docx}});
 assert.equal(final.status(),409);assert.equal((await final.json()).bloqueado,true);
 """)
        directory = Path(self.s.C.caminho(self.cid)) / "03-relatorios"
        self.assertEqual(len(list(directory.glob("minuta-v*.md"))), 2)
        self.assertEqual(len(list(directory.glob("*.docx"))), 2)
        from docx import Document
        for path in directory.glob("*.docx"):
            text = "\n".join(p.text for p in Document(path).paragraphs)
            self.assertIn("Rascunho fictício A" if "v01" in path.name else "Rascunho fictício B", text)
        self.assertFalse(any(directory.glob("*FINAL*")))

    def test_falha_retenta_sem_perder_e_conflito_exige_comparar(self):
        # Omitir revisão esperada permitiria sobrescrever a edição concorrente.
        self.browser(r"""
 await open();await text().fill('Texto local conservado');
 await context.request.post(base+'/_teste_ast12',{headers:{'X-CPJ':'1'},data:{fail:true}});
 await page.locator('#ed-salvar').click();await state('erro');
 assert.equal(await text().inputValue(),'Texto local conservado');
 await page.locator('#ed-salvar').click();await state('salvo');
 const other=await context.newPage();await other.goto(base);await other.waitForFunction(()=>!document.querySelector('#app').hidden);
 await other.evaluate(id=>abrirEditor(id),cid);
 await text().fill('Mudança do primeiro editor');await page.locator('#ed-salvar').click();await state('salvo');
 await other.locator('#ed-secoes textarea[data-s="RESUMO DOS FATOS"]').fill('Mudança do segundo editor');
 await other.locator('#ed-salvar').click();await other.waitForFunction(()=>document.querySelector('#editor').dataset.estado==='conflito');
 assert.equal(await other.locator('#ed-secoes textarea[data-s="RESUMO DOS FATOS"]').inputValue(),'Mudança do segundo editor');
 await other.locator('#ed-salvar').click();await other.waitForFunction(()=>document.querySelector('#editor').dataset.estado==='conflito');
 let latest=await context.request.get(`${base}/api/casos/${cid}/minuta`);assert.equal((await latest.json()).versao_num,2);
 await other.locator('#ed-comparar').click();await other.locator('#ed-servidor').waitFor();
 assert.match(await other.locator('#ed-servidor').inputValue(),/Mudança do primeiro editor/);
 await other.locator('#ed-reconciliar').click();await other.locator('#ed-salvar').click();
 await other.waitForFunction(()=>document.querySelector('#editor').dataset.estado==='salvo');
 latest=await context.request.get(`${base}/api/casos/${cid}/minuta`);
 assert.equal((await latest.json()).secoes['RESUMO DOS FATOS'],'Mudança do segundo editor');
 assert.equal(await context.storageState().then(s=>s.origins.some(o=>o.localStorage.some(v=>v.value.includes('Mudança')))),false);
 """)

    def test_protecao_fechar_navegar_logout_reload_e_resposta_atrasada(self):
        # Remover guardas ou aceitar GET antigo apagaria texto da sessão atual.
        self.browser(r"""
 await open();await text().fill('Não perder este texto fictício');
 const refuse=async(action)=>{const dialog=page.waitForEvent('dialog');const act=action();await (await dialog).dismiss();await act};
 await refuse(()=>page.locator('#ed-cancelar').click());assert.equal(await page.locator('#editor').isVisible(),true);
 await refuse(()=>page.evaluate(()=>abrirAba('pesquisa')));assert.equal(await text().inputValue(),'Não perder este texto fictício');
 await refuse(()=>page.evaluate(id=>abrirCaso(id),cid2));assert.equal(await text().inputValue(),'Não perder este texto fictício');
 await refuse(()=>page.evaluate(id=>abrirEditor(id),cid2));assert.equal(await text().inputValue(),'Não perder este texto fictício');
 await refuse(()=>page.evaluate(()=>document.querySelector('#b-sair').onclick()));
 assert.equal(await page.locator('#app').isVisible(),true);
 await refuse(()=>page.reload({timeout:3000}).catch(()=>{}));assert.equal(await text().inputValue(),'Não perder este texto fictício');
 await state('salvo');await page.locator('#ed-cancelar').click();
 await context.request.post(base+'/_teste_ast12',{headers:{'X-CPJ':'1'},data:{delay:0.7}});
 const old=page.evaluate(id=>abrirEditor(id),cid);await page.waitForTimeout(120);
 await open(cid2);await old;assert.match(await page.locator('#ed-tit').textContent(),new RegExp(cid2));
 assert.equal(await text().inputValue(),'');
 """)

    def test_digitacao_durante_salvamento_aguarda_confirmacao_sem_apagar_texto(self):
        # Marcar a resposta do snapshot antigo como salvo perderia a edição posterior.
        self.browser(r"""
 await open();await text().fill('Snapshot anterior fictício');
 await context.request.post(base+'/_teste_ast12',{headers:{'X-CPJ':'1'},data:{delay_post:0.6}});
 await page.locator('#ed-salvar').click();await state('salvando');
 await text().fill('Texto novo durante o salvamento');
 await page.locator('#ed-cancelar').click();assert.equal(await page.locator('#editor').isVisible(),true);
 await state('salvo');assert.equal(await text().inputValue(),'Texto novo durante o salvamento');
 let r=await context.request.get(`${base}/api/casos/${cid}/minuta`);let m=await r.json();
 assert.equal(m.versao_num,2);assert.equal(m.secoes['RESUMO DOS FATOS'],'Texto novo durante o salvamento');
 await page.evaluate(()=>abrirAba('pesquisa'));await page.evaluate(()=>abrirAba('casos'));await open();
 assert.equal(await text().inputValue(),'Texto novo durante o salvamento');
 """)

    def test_falha_ao_abrir_mostra_erro_e_permite_reabrir(self):
        # Um GET falho não pode manter o indicador de carregamento para sempre.
        self.browser(r"""
 await context.request.post(base+'/_teste_ast12',{headers:{'X-CPJ':'1'},data:{fail_get:true}});
 await page.evaluate(id=>abrirEditor(id),cid);await state('erro');
 assert.match(await page.locator('#ed-msg').textContent(),/Falha de leitura fictícia/);
 assert.equal(await page.locator('#ed-salvar').isDisabled(),true);
 await page.locator('#ed-cancelar').click();await open();
 await text().fill('Rascunho recuperado');await page.locator('#ed-salvar').click();await state('salvo');
 """)


if __name__ == "__main__":
    unittest.main(verbosity=2)
