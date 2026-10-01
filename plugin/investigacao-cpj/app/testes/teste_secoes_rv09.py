#!/usr/bin/env python3
"""RV09/V11: gerar_docx.py aceita seções com ## ou ###, títulos sem acento/caixa, acusa seção obrigatória ausente com erro
claro e acha o modelo pelo workspace (nunca por caminho fixo). Workspace temporário; o modelo DOCX é só copiado."""
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
from teste_gate_v04 import LIMPA  # noqa: E402

SCRIPT = AQUI.parents[1] / "skills" / "relatorio-ip-fraude" / "scripts" / "gerar_docx.py"
RAIZ = AQUI.parents[3]
MODELO = RAIZ / "modelos" / "MODELO RELATORIO DE INVESTIGACAO - CPJ 2026.docx"


@unittest.skipUnless(MODELO.is_file(), "modelo DOCX oficial não está neste workspace")
class SecoesRV09(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="cpj-secoes-rv09-", ignore_cleanup_errors=True)
        cls.ws = Path(cls.temp.name)
        (cls.ws / "casos").mkdir()
        (cls.ws / "modelos").mkdir()
        shutil.copy(MODELO, cls.ws / "modelos" / MODELO.name)
        for nome in ("brasao_pcsp.png", "dados-padrao.json"):
            if (RAIZ / "modelos" / nome).is_file():
                shutil.copy(RAIZ / "modelos" / nome, cls.ws / "modelos" / nome)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def gerar(self, minuta, extra=()):
        md = self.ws / "minuta.md"
        md.write_text(minuta, encoding="utf-8")
        saida = self.ws / "saida.docx"
        if saida.exists():
            saida.unlink()
        env = dict(os.environ, CPJ_WORKSPACE=str(self.ws), PYTHONIOENCODING="utf-8")
        r = subprocess.run([sys.executable, str(SCRIPT), str(md), "--saida", str(saida), *extra], cwd=str(self.ws), env=env,
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        return r, saida

    @staticmethod
    def texto(docx_path):
        import docx
        d = docx.Document(str(docx_path))
        return "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)

    def test_01_minuta_padrao_gera_docx_com_o_conteudo(self):
        r, saida = self.gerar(LIMPA)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        t = self.texto(saida)
        self.assertIn("comprovante de transferência de R$ 1.500,00", t)
        self.assertIn("Há indícios de que os valores transferidos", t)

    def test_02_secoes_com_tres_hashes_e_sem_acento_nao_perdem_conteudo(self):
        m = (LIMPA.replace("## RESUMO DOS FATOS", "### Resumo dos Fatos")
                  .replace("## DILIGÊNCIAS REALIZADAS", "## DILIGENCIAS REALIZADAS")
                  .replace("## CONCLUSÃO", "### conclusao"))
        r, saida = self.gerar(m)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        t = self.texto(saida)
        self.assertIn("induzida a realizar transferências via Pix", t, "resumo (###) preservado")
        self.assertIn("comprovante de transferência de R$ 1.500,00", t, "diligências (sem acento) preservadas")
        self.assertIn("sem elementos, até o momento, para afirmar a autoria", t, "conclusão (### minúsculo) preservada")

    def test_03_secao_obrigatoria_ausente_e_erro_claro(self):
        ini = LIMPA.index("## CONCLUSÃO")
        r, saida = self.gerar(LIMPA[:ini])
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("CONCLUSÃO", r.stdout + r.stderr)
        self.assertIn("obrigatória", (r.stdout + r.stderr).lower())
        self.assertFalse(saida.exists(), "nenhum DOCX incompleto é gerado")

    def test_04_secao_vazia_tambem_e_erro(self):
        m = LIMPA[:LIMPA.index("## CONCLUSÃO")] + "## CONCLUSÃO\n\n"
        r, _ = self.gerar(m)
        self.assertNotEqual(r.returncode, 0)

    def test_05_modelo_resolvido_pelo_workspace_sem_caminho_fixo(self):
        src = SCRIPT.read_text(encoding="utf-8")
        self.assertNotIn(r"C:\CPJ", src)
        self.assertNotIn("C:/CPJ", src)
        r, saida = self.gerar(LIMPA)  # cwd e workspace são a pasta temporária
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertTrue(saida.exists())


if __name__ == "__main__":
    unittest.main(verbosity=2)
