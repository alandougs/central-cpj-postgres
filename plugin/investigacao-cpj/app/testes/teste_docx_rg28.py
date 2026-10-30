"""RG28: regras vigentes de folhas e bloco de assinatura, só modelo geométrico."""
import hashlib
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from fixtures_docx import criar_modelo

GERADOR = Path(__file__).resolve().parents[2] / "skills" / "relatorio-ip-fraude" / "scripts" / "gerar_docx.py"


class DocxRG28(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="cpj-docx-rg28-")
        self.addCleanup(self.tmp.cleanup)
        self.ws = Path(self.tmp.name)
        self.modelo = criar_modelo(self.ws / "modelos")
        self.minuta = self.ws / "relatorio.md"
        self.saida = self.ws / "relatorio.docx"

    def gerar(self, resumo, diligencias="Exame sintético de documentos."):
        self.minuta.write_text("---\nordem_servico: FICTÍCIA/2099\ndelegado: DELEGADO FICTÍCIO\ndelegado_genero: M\n---\n"
                              "## RESUMO DOS FATOS\n" + resumo + "\n## DILIGÊNCIAS REALIZADAS\n" + diligencias +
                              "\n## CONCLUSÃO\nApenas conteúdo sintético para teste.\n", encoding="utf-8")
        fontes = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (self.modelo, self.minuta)}
        proc = subprocess.run([sys.executable, str(GERADOR), str(self.minuta), "--saida", str(self.saida),
                               "--modelo", str(self.modelo), "--sem-fluxograma"], cwd=self.ws,
                              env={**os.environ, "CPJ_WORKSPACE": str(self.ws), "CPJ_SEM_AGENTE_EMBUTIDO": "1", "PYTHONUTF8": "1"},
                              capture_output=True, text=True, encoding="utf-8", timeout=30)
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        self.assertEqual(fontes, {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in fontes})
        return Document(self.saida)

    def test_folhas_documentadas_preservadas_em_texto_e_tabela(self):
        doc = self.gerar("Origem (págs. 242–243 do PDF; fls. 12-13). Documento (Fls. 27).",
                         "| Operação | Fonte |\n|---|---|\n| Sintética | (pág. 20 do PDF; fls. 30-31) |")
        texto = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        for ref in ("(Fls. 12-13)", "(Fls. 27)", "(Fls. 30-31)"): self.assertIn(ref, texto)
        self.assertNotIn("do PDF", texto)
        self.assertNotIn("Fls. 242", texto)

    def test_pagina_sem_folha_pede_dado_sem_inventar_numero(self):
        doc = self.gerar("Origem (pág. 19 do PDF), sem folha documentada.")
        texto = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("[OBTER: FOLHAS DOS AUTOS]", texto)
        self.assertNotIn("Fls. 19", texto)
        self.assertNotIn("do PDF", texto)
        alerta = [r for p in doc.paragraphs for r in p.runs if "[OBTER: FOLHAS DOS AUTOS]" in r.text]
        self.assertTrue(alerta)
        self.assertTrue(all(str(r.font.color.rgb) == "FF0000" for r in alerta))

    def test_assinatura_tem_largura_definida_sem_alterar_proporcao_ou_imagem(self):
        doc = self.gerar("Relato sintético (Fls. 14).")
        shape = doc.inline_shapes[0]
        self.assertEqual(shape.width, Cm(4.5))
        self.assertAlmostEqual(shape.height / shape.width, 0.2, places=5)
        original = Document(self.modelo).inline_shapes[0]
        rid = shape._inline.xpath(".//a:blip")[0].get(qn("r:embed"))
        rid_original = original._inline.xpath(".//a:blip")[0].get(qn("r:embed"))
        self.assertEqual(doc.part.related_parts[rid].blob, Document(self.modelo).part.related_parts[rid_original].blob)
        self.assertEqual(shape._inline.xpath(".//a:srcRect"), [])

    def test_folha_duvidosa_preserva_interrogacao_sem_completar_digito(self):
        doc = self.gerar("Documento (PÁGINAS 8 a 9 do PDF; fl. 1? [dígito incerto]). Outro (fls. 41-42).")
        texto = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("(Fls. 1? [dígito incerto])", texto)
        self.assertIn("(Fls. 41-42)", texto)
        self.assertNotIn("do PDF", texto)

    def test_referencia_so_pdf_na_tabela_exige_folha_sem_inventar_numero(self):
        doc = self.gerar("Relato sintético (Fls. 14).",
                         "| Operação | Fonte |\n|---|---|\n| Sintética | (págs. 8-9 do PDF) |")
        fonte = doc.tables[0].cell(1, 1)
        self.assertEqual(fonte.text, "[OBTER: FOLHAS DOS AUTOS]")
        self.assertTrue(all(str(r.font.color.rgb) == "FF0000" for r in fonte.paragraphs[0].runs if r.text))


if __name__ == "__main__": unittest.main(verbosity=2)
