"""Fila de O.S. (ferramentas/fila-os.py) e reaproveitamento de extração entre casos (extrair.py).
Somente dados fictícios em pasta temporária. Uso: python teste_fila_os_claude.py"""
import os, subprocess, sys, tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[4]
FILA = RAIZ / "ferramentas" / "fila-os.py"
EXTRAIR = RAIZ / "plugin" / "investigacao-cpj" / "skills" / "pdf-autos-policiais" / "scripts" / "extrair.py"
falhas = 0


def ok(cond, nome):
    global falhas
    print(("  OK    " if cond else "  FALHA ") + nome)
    falhas += not cond


def roda(script, *args, env=None):
    r = subprocess.run([sys.executable, str(script), *args], capture_output=True, text=True, encoding="utf-8",
                       env={**os.environ, "PYTHONIOENCODING": "utf-8", **(env or {})})
    return r.returncode, r.stdout + r.stderr


with tempfile.TemporaryDirectory() as tmp:
    pasta = Path(tmp, "ordens")
    for nome, arquivos in {"OS 9001-26 - X": ["a.pdf"], "os 9002.2026": ["b.pdf"],
                           "OS 9003-2026 - Y": ["c.pdf", "Relatorio de Investigacao - OS 9003.docx"]}.items():
        Path(pasta, nome).mkdir(parents=True)
        for a in arquivos: Path(pasta, nome, a).write_bytes(b"x")
    env = {"CPJ_PASTA_OS": str(pasta)}
    print("Fila de O.S.")
    c, s = roda(FILA, "listar", env=env)
    ok(c == 0 and "9001/2026: LIVRE" in s and "9003/2026: COM_RELATORIO" in s, "detecta O.S. livres e com relatório pelo nome da pasta")
    c, s = roda(FILA, "proxima", "--agente", "A", env=env)
    ok(c == 0 and "9001" in s, "proxima assume a primeira livre")
    c, s = roda(FILA, "assumir", "9001", "--agente", "B", env=env)
    ok(c != 0 and "em andamento por A" in s, "recusa O.S. reservada por outro agente")
    c, s = roda(FILA, "assumir", "9003", "--agente", "B", env=env)
    ok(c != 0 and "já tem relatório" in s, "recusa O.S. com relatório sem --forcar")
    c, s = roda(FILA, "concluir", "9001", "--agente", "B", env=env)
    ok(c != 0, "só o dono conclui")
    import docx
    ruim, bom = Path(tmp, "ruim.docx"), Path(tmp, "bom.docx")
    d = docx.Document(); d.add_paragraph("EXCELENTÍSSIMO (A) SENHOR (A)"); d.save(ruim)
    d = docx.Document(); d.add_paragraph("EXCELENTÍSSIMO SENHOR DOUTOR DELEGADO DE POLÍCIA, o investigado(a)"); d.save(bom)
    c, s = roda(FILA, "concluir", "9001", "--agente", "A", "--docx", str(ruim), env=env)
    ok(c != 0 and "NÃO CONFERE" in s, "recusa concluir com DOCX que tem texto-guia do modelo")
    c, s = roda(FILA, "concluir", "9001", "--agente", "A", "--docx", str(bom), env=env)
    ok(c == 0 and "concluida" in s, "dono conclui (DOCX limpo; 'investigado(a)' não é resíduo)")
    ok(Path(pasta, "_CONTROLE-OS.md").exists() and "SITUAÇÃO: CONCLUIDA" in Path(pasta, "OS 9001-26 - X", "_STATUS-OS.txt").read_text(encoding="utf-8"),
       "quadro e aviso por pasta gravados")
    c, s = roda(FILA, "proxima", "--agente", "C", env=env)
    ok(c == 0 and "9002" in s, "proxima pula concluída e com relatório")

    print("Reaproveitamento de extração entre casos")
    from PIL import Image, ImageDraw
    im = Image.new("RGB", (1200, 800), "white"); ImageDraw.Draw(im).text((50, 50), "PAGINA FICTICIA R$ 1.234,56", fill="black")
    pdf = Path(tmp, "t.pdf"); im.save(pdf)
    ws = Path(tmp, "ws")
    env2 = {"CPJ_WORKSPACES": str(ws)}
    c, s = roda(EXTRAIR, str(pdf), "--saida", str(ws / "casos" / "OS-1-2026" / "01-extracao" / "t"), env=env2)
    ok(c == 0 and "reaproveitando" not in s, "primeiro caso extrai normalmente")
    c, s = roda(EXTRAIR, str(pdf), "--saida", str(ws / "casos" / "OS-2-2026" / "01-extracao" / "t"), env=env2)
    ok(c == 0 and "reaproveitando" in s and '"paginas_reaproveitadas": 1' in s, "mesmo PDF em outro caso reaproveita sem novo OCR")
    c, s = roda(EXTRAIR, str(pdf), "--saida", str(ws / "casos" / "OS-3-2026" / "01-extracao" / "t"), "--sem-reaproveitar", env=env2)
    ok(c == 0 and "reaproveitando" not in s, "--sem-reaproveitar desliga")

print("\nTUDO OK" if not falhas else f"\n{falhas} FALHA(S)")
sys.exit(1 if falhas else 0)
