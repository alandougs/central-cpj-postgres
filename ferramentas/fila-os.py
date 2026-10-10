#!/usr/bin/env python
"""Fila de Ordens de Serviço: um agente por O.S., visível a Claude, Codex, Gemini/Antigravity.

Evita que dois agentes façam o mesmo relatório. O registro fica na própria pasta das O.S.
(padrão E:\\ORDENS DE SERVIÇO CPJ), que todos os agentes enxergam, e cada pasta de O.S.
recebe um aviso _STATUS-OS.txt legível por quem só lista diretórios.

python ferramentas/fila-os.py listar                      # O.S., relatórios achados, reservas
python ferramentas/fila-os.py proxima --agente Codex-1     # assume a próxima O.S. livre e sem relatório
python ferramentas/fila-os.py assumir 2716 --agente Claude-A
python ferramentas/fila-os.py renovar 2716 --agente Claude-A      # sinal de vida (reserva vence em 4 h)
python ferramentas/fila-os.py concluir 2716 --agente Claude-A --docx "C:\\...\\RELATORIO.docx"
python ferramentas/fila-os.py liberar 2716 --agente Claude-A --motivo "parei na análise; extração pronta"
python ferramentas/fila-os.py marcar 2534 concluida --motivo "relatório V2 do investigador"   # uso do operador
"""
import argparse
from contextlib import contextmanager
import datetime
import json
import os
from pathlib import Path
import re
import sys
import tempfile

RAIZ = Path(__file__).resolve().parent.parent
_pasta_interna = RAIZ / "ordens-de-servico"
PASTA_OS = Path(os.environ.get("CPJ_PASTA_OS", _pasta_interna if _pasta_interna.is_dir() else r"E:\ORDENS DE SERVIÇO CPJ"))
# Workspaces onde agentes criam casos (o paralelo em E: existe por uso histórico).
WORKSPACES = [RAIZ] + [Path(p) for p in (r"E:\CPJ - TRABALHO",) if Path(p).resolve() != RAIZ.resolve() and Path(p).is_dir()]
_ws_env = os.environ.get("CPJ_WORKSPACES")
if _ws_env:
    WORKSPACES = [Path(p) for p in _ws_env.split(";") if p.strip()]
else:
    WORKSPACES = [RAIZ] + [Path(p) for p in (r"E:\CPJ - TRABALHO",) if Path(p).resolve() != RAIZ.resolve() and Path(p).is_dir()]
VALIDADE_H = 4
REGISTRO = "_CONTROLE-OS.json"
QUADRO = "_CONTROLE-OS.md"
AVISO = "_STATUS-OS.txt"
ESTADOS = ("livre", "em_andamento", "concluida", "suspensa")


def agora():
    return datetime.datetime.now().replace(microsecond=0)


def numeros_os(nome):
    """Extrai todos os números de O.S. e anos presentes no nome da pasta.
    Suporta separadores usuais (-, ., /, _) e múltiplas O.S. (ex.: 'OS 9999_99 e OS 9998_99 IP 8888_99 PESSOA FICTICIA').
    Desconsidera números explicitamente rotulados como IP, Processo ou BO."""
    achados = []
    # Remove menções a IP, Processo ou BO para não confundi-los com número de O.S.
    texto = re.sub(r"\b(?:IP|BO|PROC|PROCESSO|INQU[ÉE]RITO)\s*n?[º°o.]*\s*\d{1,7}\s*[-./_]\s*\d{2,4}\b", " ", nome, flags=re.I)

    # 1. Tenta identificar O.S. explicitamente identificadas com prefixo OS/O.S.
    for m in re.finditer(r"\b(?:OS|O\.S\.|ORDEM(?:\s+DE\s+SERVI[ÇC]O)?)\s*n?[º°o.]*\s*(\d{3,5})\s*[-./_]\s*(\d{2,4})\b", texto, re.I):
        ano = m.group(2)
        achados.append((m.group(1).lstrip("0") or "0", "20" + ano if len(ano) == 2 else ano))

    # 2. Se não houver prefixo OS explícito, busca o padrão geral número/ano
    if not achados:
        for m in re.finditer(r"\b(\d{3,5})\s*[-./_]\s*(\d{2,4})\b", texto):
            ano = m.group(2)
            achados.append((m.group(1).lstrip("0") or "0", "20" + ano if len(ano) == 2 else ano))

    vistos = set()
    unicos = []
    for par in achados:
        if par not in vistos:
            vistos.add(par)
            unicos.append(par)
    return unicos


def numero_os(nome):
    achados = numeros_os(nome)
    return achados[0] if achados else None


def relatorios_na_pasta(pasta):
    achados = []
    for p in pasta.rglob("*.docx"):
        rel = p.relative_to(pasta)
        if p.name.startswith("~$") or any(parte.lower() in ("trabalho", "_trabalho", "demo", "saidas") for parte in rel.parts[:-1]):
            continue
        if re.search(r"relat", p.name, re.I): achados.append(p)
    return achados


def casos_do_numero(num, ano):
    saida = []
    for ws in WORKSPACES:
        for pasta in (ws / "casos").glob(f"OS-*{num}-{ano}"):
            status = "?"
            try: status = json.loads((pasta / "caso.json").read_text(encoding="utf-8")).get("status", "?")
            except Exception: pass
            docx = sorted((pasta / "03-relatorios").glob("*.docx"), key=lambda p: p.stat().st_mtime)
            saida.append({"pasta": str(pasta), "status": status, "docx": [str(d) for d in docx]})
    return saida


def inventario():
    itens = {}
    if not PASTA_OS.is_dir(): raise SystemExit(f"Pasta das O.S. não encontrada: {PASTA_OS} (defina CPJ_PASTA_OS)")
    for pasta in sorted(p for p in PASTA_OS.iterdir() if p.is_dir() and not p.name.startswith(".")):
        chaves = numeros_os(pasta.name)
        if not chaves: continue
        pdfs_pasta = [str(p) for p in pasta.iterdir() if p.suffix.lower() == ".pdf" and "ocred" not in p.name.lower()
                      and not re.search(r"relat", p.name, re.I)]
        rels_pasta = [str(p) for p in relatorios_na_pasta(pasta)]
        for num, ano in chaves:
            item = itens.setdefault(num, {"os": num, "ano": ano, "pastas": [], "pdfs": [], "relatorios": []})
            if str(pasta) not in item["pastas"]:
                item["pastas"].append(str(pasta))
            for pdf in pdfs_pasta:
                if pdf not in item["pdfs"]:
                    item["pdfs"].append(pdf)
            for rel in rels_pasta:
                if rel not in item["relatorios"]:
                    item["relatorios"].append(rel)
    for item in itens.values():
        item["casos"] = casos_do_numero(item["os"], item["ano"])
    return itens


@contextmanager
def trava(agente):
    lock = PASTA_OS / (REGISTRO + ".lock")
    try:
        fd = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    except FileExistsError:
        idade = agora().timestamp() - lock.stat().st_mtime
        if idade < 120:
            raise SystemExit(f"Registro em atualização por outro agente ({lock}). Tente de novo em instantes.")
        lock.unlink()  # trava abandonada (> 2 min)
        fd = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f: f.write(f"pid={os.getpid()} agente={agente}\n")
        yield
    finally:
        lock.unlink(missing_ok=True)


def ler_registro():
    arq = PASTA_OS / REGISTRO
    if not arq.exists(): return {}
    return json.loads(arq.read_text(encoding="utf-8"))


def gravar_json(arq, dados):
    fd, tmp = tempfile.mkstemp(dir=arq.parent, prefix=".tmp-", suffix=".json")
    with os.fdopen(fd, "w", encoding="utf-8") as f: json.dump(dados, f, ensure_ascii=False, indent=2)
    os.replace(tmp, arq)


def vencida(r):
    if r.get("estado") != "em_andamento": return False
    return agora() > datetime.datetime.fromisoformat(r["renovada_em"]) + datetime.timedelta(hours=VALIDADE_H)


def situacao(item, r):
    """Estado efetivo: reserva ativa > registro concluído > relatório achado > livre."""
    if r.get("estado") == "em_andamento" and not vencida(r): return "em_andamento"
    if r.get("estado") in ("concluida", "suspensa"): return r["estado"]
    if item["relatorios"] or any(c["status"] in ("minuta", "entregue", "arquivado") and c["docx"] for c in item["casos"]):
        return "com_relatorio"
    return "livre"


def escrever_avisos(itens, reg):
    linhas = [f"# Controle das Ordens de Serviço", "",
              f"Gerado por `ferramentas\\fila-os.py` em {agora():%d/%m/%Y %H:%M}. **Antes de começar uma O.S., assuma-a** "
              f"(`python \"{RAIZ}\\ferramentas\\fila-os.py\" assumir <nº> --agente <nome>`). Não trabalhe em O.S. "
              f"`em_andamento` de outro agente nem refaça O.S. `concluida`/`com_relatorio` sem pedido do investigador.", "",
              "| O.S. | Situação | Agente | Desde | Relatório / observação |", "|---|---|---|---|---|"]
    for num in sorted(itens, key=int):
        item, r = itens[num], reg.get(num, {})
        sit = situacao(item, r)
        rel = r.get("docx") or (item["relatorios"][-1] if item["relatorios"] else "")
        if not rel:
            docs = [d for c in item["casos"] for d in c["docx"]]
            rel = docs[-1] if docs else ""
        obs = " — ".join(x for x in (Path(rel).name if rel else "", r.get("motivo", "")) if x)
        desde = r.get("renovada_em", "")[:16].replace("T", " ")
        linhas.append(f"| {num}/{item['ano']} | {sit}{' (VENCIDA)' if vencida(r) else ''} | {r.get('agente', '')} | {desde} | {obs} |")
        texto = (f"O.S. {num}/{item['ano']} — SITUAÇÃO: {sit.upper()}\n"
                 f"Agente: {r.get('agente', '-')}   Atualizado: {desde or '-'}\n"
                 f"{obs}\n\nConsulte/assuma pelo script: python \"{RAIZ}\\ferramentas\\fila-os.py\" listar\n")
        for pasta in item["pastas"]:
            try: Path(pasta, AVISO).write_text(texto, encoding="utf-8")
            except OSError: pass
    (PASTA_OS / QUADRO).write_text("\n".join(linhas) + "\n", encoding="utf-8")


def conferir_docx(caminho):
    """Resíduos do modelo que não podem chegar ao delegado (mesma regra do gerar_docx.py)."""
    if not Path(caminho).is_file(): return [f"arquivo não encontrado: {caminho}"]
    try:
        import docx
        d = docx.Document(caminho)
    except Exception as e:
        return [f"DOCX ilegível ({e})"]
    texto = "\n".join([p.text for p in d.paragraphs] + [q.text for t in d.tables for c in t._cells for q in c.paragraphs])
    return sorted({m.group(0) for m in re.finditer(r" \((?:A|a)\)|A\(o\)|\{[^}\n]{0,40}\}?|\[n[º°o][^\]\n]{0,20}\]?|\bPREENCHER\b", texto)})


def alterar(acao, num, agente, docx=None, motivo=None, estado=None, forcar=False):
    if acao == "concluir" and docx and not forcar:
        problemas = conferir_docx(docx)
        if problemas:
            raise SystemExit("DOCX NÃO CONFERE — corrija antes de concluir (ou --forcar): " + " | ".join(problemas[:10]))
    itens = inventario()
    with trava(agente or "operador"):
        reg = ler_registro()
        if acao == "proxima":
            livres = [n for n in sorted(itens, key=int) if situacao(itens[n], reg.get(n, {})) == "livre" and itens[n]["pdfs"]]
            if not livres: raise SystemExit("Nenhuma O.S. livre e sem relatório.")
            num, acao = livres[0], "assumir"
        if num not in itens: raise SystemExit(f"O.S. {num} não encontrada em {PASTA_OS}.")
        r = reg.get(num, {})
        sit = situacao(itens[num], r)
        if acao == "assumir":
            if sit == "em_andamento" and r.get("agente") != agente:
                raise SystemExit(f"O.S. {num} em andamento por {r['agente']} desde {r['assumida_em']}. Escolha outra.")
            if sit in ("concluida", "com_relatorio") and not forcar:
                raise SystemExit(f"O.S. {num} já tem relatório ({sit}). Use --forcar só com pedido do investigador.")
            r = {"estado": "em_andamento", "agente": agente, "assumida_em": agora().isoformat(), "renovada_em": agora().isoformat()}
        elif acao in ("renovar", "concluir", "liberar"):
            if r.get("agente") != agente or r.get("estado") != "em_andamento":
                raise SystemExit(f"O.S. {num} não está reservada para {agente} (estado: {r.get('estado', 'livre')}, agente: {r.get('agente', '-')}).")
            r["renovada_em"] = agora().isoformat()
            if acao == "concluir": r.update(estado="concluida", docx=docx or "")
            if acao == "liberar": r.update(estado="livre")
            if motivo: r["motivo"] = motivo
        elif acao == "marcar":
            r = {**r, "estado": estado, "agente": r.get("agente", "operador"), "renovada_em": agora().isoformat(),
                 **({"motivo": motivo} if motivo else {}), **({"docx": docx} if docx else {})}
        r.setdefault("historico", [])
        r["historico"] = (r.get("historico") or []) + [f"{agora().isoformat()} {acao} {agente or 'operador'}"]
        reg[num] = r
        gravar_json(PASTA_OS / REGISTRO, reg)
        escrever_avisos(itens, reg)
    print(f"O.S. {num}: {acao} ok ({r['estado']}, {r.get('agente', '')}).")


def listar(como_json=False):
    itens, reg = inventario(), ler_registro()
    if como_json:
        print(json.dumps({n: {**i, "registro": reg.get(n, {}), "situacao": situacao(i, reg.get(n, {}))} for n, i in itens.items()},
                         ensure_ascii=False, indent=2)); return
    for num in sorted(itens, key=int):
        item, r = itens[num], reg.get(num, {})
        print(f"O.S. {num}/{item['ano']}: {situacao(item, r).upper()}" + (f"  [{r.get('agente')}]" if r.get("agente") else ""))
        for p in item["pdfs"]: print(f"   PDF  {p}")
        for p in item["relatorios"]: print(f"   DOCX {p}")
        for c in item["casos"]: print(f"   CASO {c['pasta']}  status={c['status']}  docx={len(c['docx'])}")
    with trava("listar"):
        escrever_avisos(itens, reg)


def main():
    for fluxo in (sys.stdout, sys.stderr):
        if hasattr(fluxo, "reconfigure"): fluxo.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    l = sub.add_parser("listar"); l.add_argument("--json", action="store_true")
    p = sub.add_parser("proxima"); p.add_argument("--agente", required=True)
    for nome in ("assumir", "renovar", "concluir", "liberar"):
        s = sub.add_parser(nome); s.add_argument("os"); s.add_argument("--agente", required=True)
        s.add_argument("--motivo"); s.add_argument("--docx"); s.add_argument("--forcar", action="store_true")
    m = sub.add_parser("marcar"); m.add_argument("os"); m.add_argument("estado", choices=ESTADOS)
    m.add_argument("--motivo"); m.add_argument("--docx")
    a = ap.parse_args()
    if a.cmd == "listar": return listar(a.json)
    if a.cmd == "proxima": return alterar("proxima", None, a.agente)
    if a.cmd == "marcar": return alterar("marcar", a.os.lstrip("0"), None, a.docx, a.motivo, a.estado)
    alterar(a.cmd, a.os.lstrip("0"), a.agente, a.docx, a.motivo, forcar=a.forcar)


if __name__ == "__main__":
    main()
