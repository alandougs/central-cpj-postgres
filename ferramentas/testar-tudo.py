#!/usr/bin/env python
"""Executa as suítes do app em sandboxes e a suíte HTTP da Central."""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import secrets
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request


ROOT = Path(__file__).resolve().parent.parent
PLUGIN = ROOT / "plugin" / "investigacao-cpj"
TESTES = PLUGIN / "app" / "testes"
REGRESSAO_RAPIDA = (
    "teste_solo_claude.py",
    "teste_seguranca_codex.py",
    "teste_modularizacao_gemini.py",
    "teste_dados_os_codex.py",
)


def argumentos():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--rapido", action="store_true", help="roda a regressão mínima")
    ap.add_argument("--incluir", metavar="TESTE", help="inclui teste_*.py na execução rápida")
    return ap.parse_args()


def comando_python(script: Path) -> list[str]:
    return [sys.executable, "-X", "utf8", "-W", "ignore::ResourceWarning", str(script)]


def clonar_codigo(destino: Path) -> Path:
    """Cria snapshot descartável para que testes não escrevam no workspace real."""
    raiz_teste = destino / "codigo"
    (raiz_teste / "plugin").mkdir(parents=True)
    shutil.copytree(PLUGIN, raiz_teste / "plugin" / "investigacao-cpj")
    shutil.copytree(ROOT / "modelos", raiz_teste / "modelos")
    shutil.copytree(ROOT / "ferramentas", raiz_teste / "ferramentas",
                    ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
    shutil.copytree(ROOT / "portatil", raiz_teste / "portatil")
    agentes = raiz_teste / ".agents" / "skills"
    agentes.mkdir(parents=True)
    for skill in (ROOT / ".agents" / "skills").glob("cpj-*"):
        destino_skill = agentes / skill.name
        if skill.is_dir():
            shutil.copytree(skill, destino_skill)
        else:
            shutil.copy2(skill, destino_skill)
    for nome in ("AGENTS.md", "CLAUDE.md", "GEMINI.md", "PRD.md", "TAREFAS-COMPARTILHADAS.md", "Central CPJ.bat"):
        origem = ROOT / nome
        if origem.is_file():
            shutil.copy2(origem, raiz_teste / nome)
    for nome in ("requirements.txt", "requirements-opcional.txt", "requirements-dev.txt"):
        origem = ROOT / nome
        if origem.is_file():
            shutil.copy2(origem, raiz_teste / nome)
    modelo = raiz_teste / "casos" / "_MODELO-CASO"
    for pasta in ("00-originais", "01-extracao", "02-analise", "03-relatorios"):
        (modelo / pasta).mkdir(parents=True, exist_ok=True)
    (modelo / "caso.json").write_text(json.dumps({
        "id": "_MODELO-CASO", "status": "recebido", "datas": {}, "relatorios": []
    }, ensure_ascii=False), encoding="utf-8")
    (raiz_teste / "revisoes").mkdir()
    return raiz_teste


def validar_nome_teste(nome: str) -> str:
    if Path(nome).name != nome or not nome.startswith("teste_") or not nome.endswith(".py"):
        raise ValueError("--incluir deve ser o nome de um arquivo teste_*.py")
    if not (TESTES / nome).is_file():
        raise ValueError(f"teste não encontrado: {nome}")
    return nome


def selecionar_testes(rapido: bool, incluir: str | None) -> list[str]:
    todos = sorted(p.name for p in TESTES.glob("teste_*.py") if p.name != "teste_central.py")
    if not rapido:
        return todos
    nomes = list(REGRESSAO_RAPIDA)
    if incluir:
        nomes.append(validar_nome_teste(incluir))
    ausentes = [nome for nome in nomes if nome not in todos]
    if ausentes:
        raise ValueError("testes da regressão não encontrados: " + ", ".join(ausentes))
    return list(dict.fromkeys(nomes))


def preparar_workspace_teste(raiz_teste: Path, codigo: Path) -> Path:
    ws = raiz_teste / "workspace"
    (ws / "config").mkdir(parents=True)
    shutil.copytree(ROOT / "modelos", ws / "modelos")
    shutil.copytree(codigo / "casos" / "_MODELO-CASO", ws / "casos" / "_MODELO-CASO")
    return ws


def rodar_suíte(nome: str) -> tuple[int, float, str]:
    with tempfile.TemporaryDirectory(prefix=f"cpj-suite-{Path(nome).stem}-") as td:
        raiz_teste = Path(td)
        codigo = clonar_codigo(raiz_teste)
        ws = preparar_workspace_teste(raiz_teste, codigo)
        script = codigo / "plugin" / "investigacao-cpj" / "app" / "testes" / nome
        env = os.environ.copy()
        env["CPJ_WORKSPACE"] = str(ws)
        env["CPJ_SEM_AGENTE_EMBUTIDO"] = "1"
        inicio = time.monotonic()
        try:
            proc = subprocess.run(comando_python(script), cwd=raiz_teste, env=env,
                                  capture_output=True, text=True, encoding="utf-8", errors="replace",
                                  timeout=600)
        except subprocess.TimeoutExpired as e:
            saida = e.stdout or ""
            erro = e.stderr or ""
            if isinstance(saida, bytes):
                saida = saida.decode("utf-8", errors="replace")
            if isinstance(erro, bytes):
                erro = erro.decode("utf-8", errors="replace")
            return 124, 600.0, saida + erro + "\nSuíte excedeu o limite de 600 segundos."
        return proc.returncode, time.monotonic() - inicio, proc.stdout + proc.stderr


def porta_livre() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def gerar_fixtures(ws: Path) -> None:
    import openpyxl
    from docx import Document
    from pypdf import PdfReader, PdfWriter

    env = os.environ.copy()
    pdf_script = PLUGIN / "skills" / "pdf-autos-policiais" / "teste" / "gerar_pdf_ficticio.py"
    pdf = subprocess.run(comando_python(pdf_script), cwd=ws, env=env, capture_output=True,
                         text=True, encoding="utf-8", errors="replace", timeout=180)
    if pdf.returncode:
        raise RuntimeError("gerador do PDF fictício falhou:\n" + pdf.stdout + pdf.stderr)
    produzido = ws / "ip_teste.pdf"
    if not produzido.is_file():
        raise RuntimeError("gerador não criou ip_teste.pdf")
    leitor = PdfReader(str(produzido))
    writer = PdfWriter()
    import io
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    cv = canvas.Canvas(buf, pagesize=A4)
    cv.setFont("Helvetica", 12)
    cv.drawString(72, 800, "INQUERITO POLICIAL FICTICIO - fls. 1")
    cv.drawString(72, 770, "Termo de declaracoes de FULANO FICTICIO DE TAL. CPF 999.888.777-66.")
    cv.drawString(72, 750, "Mae: BELTRANA FICTICIA. Telefone: (18) 99777-1122.")
    cv.drawString(72, 730, "Transferencia realizada para a chave Pix 99988877766.")
    cv.showPage()
    cv.save()
    p_extra = PdfReader(io.BytesIO(buf.getvalue())).pages[0]
    writer.add_page(p_extra)
    for indice in (0, 80):  # página textual e uma imagem OCR
        if indice < len(leitor.pages):
            writer.add_page(leitor.pages[indice])
    with (ws / "ip_ficticio.pdf").open("wb") as f:
        writer.write(f)
    produzido.unlink()

    wb = openpyxl.Workbook()
    sh = wb.active
    sh.append(["Nome", "CPF", "Nome da mãe", "Telefone"])
    sh.append(["FULANO FICTICIO DE TAL", "999.888.777-66", "BELTRANA FICTICIA", "99777-1122"])
    sh.append(["OUTRA PESSOA FICTICIA", "888.777.666-55", "BELTRANA FICTICIA", "99777-3344"])
    wb.save(ws / "muralha_ficticio.xlsx")

    doc = Document()
    doc.add_heading("Relatório de investigação fictício", 0)
    doc.add_paragraph("Documento exclusivamente fictício para teste automatizado.")
    doc.save(ws / "referencia_ficticia.docx")

    cred = {
        perfil: {"login": f"{perfil}_teste", "senha": secrets.token_urlsafe(24)}
        for perfil in ("admin", "delegado", "escrivao")
    }
    (ws / "config" / "credenciais-teste.json").write_text(
        json.dumps(cred, ensure_ascii=False), encoding="utf-8"
    )


def esperar_central(url: str, proc: subprocess.Popen, segundos: int = 30) -> bool:
    limite = time.monotonic() + segundos
    while time.monotonic() < limite:
        if proc.poll() is not None:
            return False
        try:
            with urllib.request.urlopen(url, timeout=1) as resposta:
                if resposta.status < 500:
                    return True
        except (OSError, urllib.error.URLError):
            time.sleep(0.25)
    return False


def rodar_central() -> tuple[int, float, str]:
    with tempfile.TemporaryDirectory(prefix="cpj-central-e2e-") as td:
        raiz_teste = Path(td)
        codigo = clonar_codigo(raiz_teste)
        ws = preparar_workspace_teste(raiz_teste, codigo)
        gerar_fixtures(ws)
        porta = porta_livre()
        url = f"http://127.0.0.1:{porta}"
        servidor = codigo / "plugin" / "investigacao-cpj" / "app" / "servidor.py"
        env = os.environ.copy()
        env["CPJ_WORKSPACE"] = str(ws)
        env["CPJ_SEM_AGENTE_EMBUTIDO"] = "1"
        log_path = raiz_teste / "servidor.log"
        inicio = time.monotonic()
        with log_path.open("w", encoding="utf-8") as log:
            proc = subprocess.Popen(comando_python(servidor) + [
                "--workspace", str(ws), "--porta", str(porta), "--somente-local", "--sem-navegador"
            ], cwd=raiz_teste, env=env, stdout=log, stderr=subprocess.STDOUT)
            try:
                if not esperar_central(url, proc):
                    log.flush()
                    return 1, time.monotonic() - inicio, "Central não iniciou.\n" + log_path.read_text(encoding="utf-8", errors="replace")
                teste = codigo / "plugin" / "investigacao-cpj" / "app" / "testes" / "teste_central.py"
                proc_teste = subprocess.run(
                    comando_python(teste) + ["--url", url, "--workspace", str(ws)],
                    cwd=raiz_teste, env=env, capture_output=True, text=True,
                    encoding="utf-8", errors="replace", timeout=600,
                )
                detalhe = proc_teste.stdout + proc_teste.stderr
                if proc_teste.returncode:
                    log.flush()
                    detalhe += "\n--- log da Central ---\n" + log_path.read_text(encoding="utf-8", errors="replace")[-12000:]
                return proc_teste.returncode, time.monotonic() - inicio, detalhe
            finally:
                proc.terminate()
                try:
                    proc.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    proc.kill()
                    proc.wait(timeout=5)


def main() -> int:
    args = argumentos()
    try:
        nomes = selecionar_testes(args.rapido, args.incluir)
    except ValueError as e:
        print(f"ERRO: {e}", file=sys.stderr)
        return 2

    resultados: list[tuple[str, int, float]] = []
    detalhes_falha: list[tuple[str, str]] = []
    for nome in nomes:
        codigo, duracao, detalhe = rodar_suíte(nome)
        resultados.append((nome, codigo, duracao))
        estado = "OK" if codigo == 0 else "FALHA"
        print(f"{estado:6} {duracao:8.2f}s  {nome}", flush=True)
        if codigo:
            detalhes_falha.append((nome, detalhe))

    if not args.rapido:
        try:
            codigo, duracao, detalhe = rodar_central()
        except Exception as e:
            codigo, duracao, detalhe = 1, 0.0, f"{type(e).__name__}: {e}"
        resultados.append(("teste_central.py (Central isolada)", codigo, duracao))
        estado = "OK" if codigo == 0 else "FALHA"
        print(f"{estado:6} {duracao:8.2f}s  teste_central.py (Central isolada)", flush=True)
        if codigo:
            detalhes_falha.append(("teste_central.py (Central isolada)", detalhe))

    print("\nResumo de tempos:")
    print(f"{'Resultado':8} {'Segundos':>9}  Suíte")
    for nome, codigo, duracao in resultados:
        print(f"{'OK' if codigo == 0 else 'FALHA':8} {duracao:9.2f}  {nome}")
    if detalhes_falha:
        print("\nFalhas:")
        for nome, detalhe in detalhes_falha:
            print(f"\n--- {nome} ---\n{detalhe[-12000:]}")
    total = sum(d for _, _, d in resultados)
    falhas = sum(1 for _, codigo, _ in resultados if codigo)
    print(f"\n{len(resultados) - falhas}/{len(resultados)} suítes aprovadas em {total:.2f}s")
    return 1 if falhas else 0


if __name__ == "__main__":
    raise SystemExit(main())
