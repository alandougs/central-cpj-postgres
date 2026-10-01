import os, sys, tempfile, json, shutil, threading, subprocess, socket, time
from pathlib import Path
ROOT = Path.cwd()
OUT = ROOT / 'revisoes/revisao-senior-2026-09-27'
W = Path(tempfile.mkdtemp(prefix='cpj-auditoria-senior-'))
os.environ['CPJ_WORKSPACE'] = str(W)
os.environ['CPJ_SEM_AGENTE_EMBUTIDO'] = '1'
shutil.copytree(ROOT/'casos/_MODELO-CASO', W/'casos/_MODELO-CASO')
(W/'config').mkdir()
cred = {p:{'login':p+'_ficticio','senha':'TesteFicticio9876'} for p in ['admin','delegado','escrivao']}
(W/'config/credenciais-teste.json').write_text(json.dumps(cred))
from reportlab.pdfgen import canvas
c = canvas.Canvas(str(W/'ip_ficticio.pdf'))
c.drawString(40,800,'DOCUMENTO INTEIRAMENTE FICTICIO PARA TESTE DE SOFTWARE')
c.drawString(40,775,'CPF 999.888.777-66. BO AB0001/2026. IP 0001/2026.')
for i in range(15): c.drawString(40,740-i*20,'Texto de teste ficticio para validar extracao de documento com texto nativo.')
c.save()
from openpyxl import Workbook
b = Workbook(); s=b.active
s.append(['Nome','Mae','CPF','Telefone'])
s.append(['FULANO FICTICIO','BELTRANA FICTICIA','99988877766','18997771122'])
s.append(['OUTRO FICTICIO','BELTRANA FICTICIA','11122233344','18998887766'])
b.save(W/'muralha_ficticio.xlsx')
from docx import Document
d=Document();d.add_paragraph('Relatorio de referencia inteiramente ficticio.');d.save(W/'referencia_ficticia.docx')
sys.path.insert(0,str(ROOT/'plugin/investigacao-cpj/app'))
import servidor as S
S.tarefas.claude_status=lambda:{'instalado':False,'logado':False,'teste':True}
threading.Thread(target=S.trabalhador,daemon=True).start()
sock=socket.socket();sock.bind(('127.0.0.1',8766));sock.close()
threading.Thread(target=lambda:S.app.run(host='127.0.0.1',port=8766,threaded=True,debug=False),daemon=True).start()
(OUT/'ambiente.json').write_text(json.dumps({'workspace':str(W),'url':'http://127.0.0.1:8766','pid':os.getpid()},indent=2))
print('SERVIDOR FICTICIO PRONTO http://127.0.0.1:8766',flush=True)
time.sleep(1)
r=subprocess.run([sys.executable,'-X','utf8',str(ROOT/'plugin/investigacao-cpj/app/testes/teste_central.py'),'--url','http://127.0.0.1:8766','--workspace',str(W)],capture_output=True,text=True,encoding='utf-8')
(OUT/'suite-central.log').write_text(r.stdout+'\n'+r.stderr,encoding='utf-8')
print('SUITE CENTRAL RETORNO',r.returncode,flush=True)
print(r.stdout,flush=True)
while not (OUT/'encerrar-servidor').exists(): time.sleep(1)
print('Servidor de teste encerrado. Fixture preservada para reproducao local.',flush=True)
